import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import verify_codes as v
from test_matching import erp_record


class RegressionTest(unittest.TestCase):
    def setUp(self):
        self.record = erp_record("000123", "暖贴艾草30片", source="单品", item_id="ITEM-A")
        self.features = v.build_features("暖贴", "艾草暖贴30片")
        self.features["erp_item_id"] = "ITEM-A"
        self.index = {"000123": [self.record]}

    def test_no_related_erp_record(self):
        result = v.classify_row("", v.build_features("无关联的天文望远镜", "专业光学镜筒"), self.index, [self.record])
        self.assertEqual(result["核验状态"], "ERP无匹配")
        self.assertEqual(result["应改为编码"], "")

    def test_multiple_identity_hits(self):
        other = erp_record("000124", "暖贴艾草30片", item_id="ITEM-A")
        result = v.classify_row("000123", self.features, self.index, [self.record, other])
        self.assertEqual(result["核验状态"], "待人工确认")
        self.assertEqual(result["独立身份命中条数"], 2)

    def test_combo_and_single_conflict(self):
        other = erp_record("000124", "暖贴艾草30片", source="组合装", item_id="ITEM-A")
        result = v.classify_row("000123", self.features, {**self.index, "000124": [other]}, [self.record, other])
        self.assertEqual(result["核验状态"], "待人工确认")
        self.assertIn("组合装", result["ERP数据来源"])
        self.assertIn("单品", result["ERP数据来源"])

    def test_global_code_conflict_blocks_different_current_code(self):
        other = erp_record("000123", "暖贴生姜30片", source="组合装", item_id="ITEM-B")
        old = erp_record("000999", "暖贴生姜30片", item_id="ITEM-C")
        result = v.classify_row("000999", self.features, {"000123": [self.record, other], "000999": [old]}, [self.record, other, old])
        self.assertEqual(result["核验状态"], "待人工确认")
        self.assertEqual(result["应改为编码"], "")

    def test_leading_zero_code_kept_when_suggesting(self):
        old = erp_record("999", "暖贴生姜30片", item_id="ITEM-B")
        result = v.classify_row("999", self.features, {**self.index, "999": [old]}, [self.record, old])
        self.assertEqual(result["应改为编码"], "000123")

    def test_missing_current_erp_code_is_exception_with_candidates(self):
        result = v.classify_row("OLD-CODE", self.features, self.index, [self.record])
        self.assertEqual(result["核验状态"], "异常（ERP未找到）")
        self.assertGreater(result["候选总数"], 0)
        self.assertEqual(result["应改为编码"], "")

    def test_missing_current_erp_code_is_exception_without_candidates(self):
        result = v.classify_row("OLD-CODE", v.build_features("无关联的天文望远镜", "专业光学镜筒"), self.index, [self.record])
        self.assertEqual(result["核验状态"], "异常（ERP未找到）")
        self.assertEqual(result["候选总数"], 0)
        self.assertEqual(result["应改为编码"], "")

    def test_blank_current_code_not_marked_erp_not_found_exception(self):
        result = v.classify_row("", self.features, self.index, [self.record])
        self.assertEqual(result["核验状态"], "编码缺失（待补录）")
        self.assertEqual(result["应改为编码"], "")

    def test_quantity_conflict_cannot_be_overridden_by_identity(self):
        self.features["qty"] = (10, "片")
        result = v.classify_row("000123", self.features, self.index, [self.record])
        self.assertEqual(result["核验状态"], "待人工确认")
        self.assertEqual(result["应改为编码"], "")

    def test_pair_conversion_with_independent_identity(self):
        record = erp_record("001", "足贴艾草10对", item_id="FOOT-A")
        features = v.build_features("足贴", "足贴艾草20片")
        features["erp_item_id"] = "FOOT-A"
        result = v.classify_row("001", features, {"001": [record]}, [record])
        self.assertEqual(result["核验状态"], "编码正确")

    def test_gift_does_not_auto_match_simple_total(self):
        self.features = v.build_features("暖贴", "艾草暖贴30片送精油")
        self.features["erp_item_id"] = "ITEM-A"
        result = v.classify_row("000123", self.features, self.index, [self.record])
        self.assertEqual(result["核验状态"], "待人工确认")

    def test_multipack_and_explicit_total_conflict(self):
        quantity, issues = v.quantity_details("5片*6盒，共20片")
        # 2026-10-28: a contradictory explicit package total is unresolved.
        self.assertIsNone(quantity)
        self.assertEqual(v.quantity_structure("5片*6盒，共20片")['parsed_total'],(20,'片'))
        self.assertTrue(issues)

    def test_box_pack_multiplier(self):
        self.assertEqual(v.extract_quantity("暖贴60片盒装*2"), (120, "片"))

    def test_three_in_one_total_excludes_each_twenty(self):
        a = v.build_features("三合一暖贴", "三合一【艾草+益母草+生姜】20片")
        b = v.build_features("", "", erp_name="艾草/益母草/生姜暖贴各20片")
        self.assertIn("规格数量不一致", v.feature_conflicts(a, b))

    def test_literal_csv_wrapper_never_evaluated(self):
        self.assertEqual(v.literal_text('="000123"'), "000123")
        self.assertEqual(v.literal_text("=1+2"), "=1+2")

    def test_alias_headers_differ(self):
        df = pd.DataFrame(columns=["itemId", "itemTitle", "skuId", "SKU规格", "skuOuterId"])
        self.assertEqual(v.choose_platform_columns(df)["code"], "skuOuterId")

    def test_duplicate_generic_code_columns_require_mapping(self):
        df = pd.DataFrame(columns=["商品ID", "商品标题", "商家编码", "商家编码.1"])
        with self.assertRaisesRegex(RuntimeError, "多个可能列"):
            v.choose_platform_columns(df)

    def test_current_code_column_cannot_be_identity(self):
        df = pd.DataFrame(columns=["商品ID", "商品标题", "商家编码"])
        v.PLATFORM_PROFILES["测试"] = {"columns": {"erp_item_id": "商家编码"}}
        try:
            with self.assertRaisesRegex(RuntimeError, "身份标识"):
                v.choose_platform_columns(df, "测试")
        finally:
            del v.PLATFORM_PROFILES["测试"]

    def test_display_limit_does_not_limit_identity_scan(self):
        records = [erp_record(f"CODE-{i}", "暖贴艾草30片", item_id="ITEM-A") for i in range(6)]
        result = v.classify_row("CODE-0", self.features, {r["code"]:[r] for r in records}, records)
        self.assertEqual(result["候选总数"], 6)
        self.assertEqual(result["独立身份命中条数"], 6)
        self.assertEqual(result["核验状态"], "待人工确认")

    def test_spec_code_disambiguates_only_explicit_identity(self):
        self.record["spec_code"] = "SPEC-A"
        other = erp_record("000124", "暖贴艾草30片", item_id="ITEM-A")
        other["spec_code"] = "SPEC-B"
        self.features["spec_code"] = "SPEC-A"
        result = v.classify_row("000123", self.features, {**self.index,"000124":[other]}, [self.record,other])
        self.assertEqual(result["核验状态"], "编码正确")


class FileWorkflowTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.combo = self.folder / "组合装.csv"
        self.single = self.folder / "单品.csv"
        pd.DataFrame(columns=["组合装名称", "商家编码"]).to_csv(self.combo, index=False, encoding="utf-8-sig")
        pd.DataFrame([{"货品名称": "暖贴艾草30片", "商家编码": '="000123"', "货品编号": "ITEM-A", "附加字段": "保留我"}]).to_csv(self.single, index=False, encoding="utf-8-sig")
        self.platform = self.folder / "custom.csv"
        self.base = {"商品ID": "00007", "skuID": "00008", "商品标题": "=2+2", "sku名称": "艾草暖贴30片", "sku编码": " 999 ", "ERP货品编号": "ITEM-A"}

    def tearDown(self):
        self.temp.cleanup()

    def run_file(self, rows):
        pd.DataFrame(rows).to_csv(self.platform, index=False, encoding="utf-8-sig")
        erp = v.load_erp(str(self.combo), str(self.single))
        records, index = v.build_erp_records(erp)
        results, meta = v.verify_platform_file(str(self.platform), records, index)
        return erp, results, meta

    def test_duplicate_sku_rows_preserved_and_reviewed(self):
        row = {**self.base, "sku编码": "000123"}
        _, result, _ = self.run_file([row, row])
        self.assertEqual(len(result), 2)
        self.assertTrue(result["核验状态"].eq("核对通过").all())
        self.assertTrue(result["应改为编码"].eq("").all())

    def test_original_values_and_excel_text_round_trip(self):
        single = pd.read_csv(self.single, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        single = pd.concat([single, pd.DataFrame([{"货品名称": "暖贴生姜30片", "商家编码": "999", "货品编号": "ITEM-B"}])], ignore_index=True).fillna("")
        single.to_csv(self.single, index=False, encoding="utf-8-sig")
        erp, result, meta = self.run_file([self.base])
        self.assertEqual(result.iloc[0]["原编码"], " 999 ")
        self.assertEqual(result.iloc[0]["应改为编码"], "")
        self.assertEqual(result.iloc[0]["核验状态"], "异常（商品/规格不一致）")
        self.assertEqual(erp.iloc[0]["附加字段"], "保留我")
        self.assertEqual(json.loads(erp.iloc[0]["原始字段JSON"])["商家编码"], '="000123"')
        output = self.folder / "report.xlsx"
        v.write_report(str(output), result, erp, [meta])
        book = load_workbook(output, data_only=False)
        sheet = book["核验结果"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for field, value in [("商品 ID", "00007"), ("SKU ID", "00008"), ("原编码", " 999 "), ("商品标题", "=2+2")]:
            cell = sheet.cell(2, headers[field])
            self.assertEqual(cell.value, value)
            self.assertEqual(cell.data_type, "s")
        self.assertIsNone(sheet.cell(2, headers['应改为编码']).value)
        self.assertEqual(len(book["平台原始数据"]["A"]), 2)
        self.assertEqual(sum(c.data_type == "f" for s in book for row in s for c in row), 0)
        book.close()

    def test_merged_same_code_and_specs_retains_both_sources(self):
        pd.DataFrame([{"组合装名称": "暖贴艾草30片", "商家编码": "000123", "货品编号": "ITEM-A"}]).to_csv(self.combo, index=False, encoding="utf-8-sig")
        erp, result, _ = self.run_file([{**self.base, "sku编码": "000123"}])
        self.assertEqual(len(erp), 2)
        self.assertEqual(set(erp["ERP来源"]), {"组合装", "单品"})
        self.assertEqual(result.iloc[0]["核验状态"], "核对通过")

    def test_blank_erp_code_retained_and_blocks_automatic(self):
        pd.DataFrame([{"货品名称": "暖贴艾草30片", "商家编码": "", "货品编号": "ITEM-A"}]).to_csv(self.single, index=False, encoding="utf-8-sig")
        erp, result, _ = self.run_file([self.base])
        self.assertEqual(len(erp), 1)
        self.assertEqual(result.iloc[0]["核验状态"], "异常（ERP未找到）")

    def test_cli_rejects_source_overwrite(self):
        pd.DataFrame([self.base]).to_csv(self.platform, index=False, encoding="utf-8-sig")
        before = self.platform.read_bytes()
        completed = subprocess.run([sys.executable, str(Path(v.__file__)), "--combo-erp", str(self.combo), "--single-erp", str(self.single), "--platform-file", str(self.platform), "--output", str(self.platform)], capture_output=True)
        self.assertNotEqual(completed.returncode, 0)
        self.assertEqual(before, self.platform.read_bytes())

    def test_excel_numeric_zero_format_is_preserved(self):
        path = self.folder / "zero.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["商品ID", "商品标题", "skuID", "sku名称", "sku编码"])
        sheet.append([7, "暖贴", 8, "艾草暖贴30片", 123])
        for address in ("A2", "C2", "E2"):
            sheet[address].number_format = "000000"
        book.save(path)
        book.close()
        frame, _, _ = v.load_platform(str(path))
        self.assertEqual(frame.iloc[0]["sku编码"], "000123")
        self.assertEqual(frame.iloc[0]["商品ID"], "000007")

    def test_csv_preamble_variable_column_count(self):
        self.platform.write_text('导出说明\n商品ID,商品标题,sku编码\n0001,暖贴,0002\n', encoding="utf-8-sig")
        frame, _, header = v.load_platform(str(self.platform))
        self.assertEqual(header, 1)
        self.assertEqual(frame.iloc[0]["sku编码"], "0002")

    def test_excel_na_identifier_is_not_missing(self):
        path = self.folder / "na.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["商品ID", "商品标题", "skuID", "sku名称", "sku编码"])
        sheet.append(["NA", "暖贴", "NULL", "艾草暖贴30片", "NA"])
        book.save(path)
        book.close()
        frame, _, _ = v.load_platform(str(path))
        self.assertEqual(frame.iloc[0]["商品ID"], "NA")
        self.assertEqual(frame.iloc[0]["sku编码"], "NA")


if __name__ == "__main__":
    unittest.main()
