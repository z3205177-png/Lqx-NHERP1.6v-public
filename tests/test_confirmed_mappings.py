import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import verify_codes as v
import confirmed_mappings as cm


class ConfirmedMappingsTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.combo, self.single = self.root/'combo.csv', self.root/'single.csv'
        self.platform = self.root/'shop.csv'
        self.report, self.mapping = self.root/'report.xlsx', self.root/'confirmed.json'
        pd.DataFrame(columns=['组合装名称','商家编码']).to_csv(self.combo,index=False,encoding='utf-8-sig')
        pd.DataFrame([{'货品名称':'艾草暖贴30片','商家编码':'A','货品编号':'ITEM-A'},
                      {'货品名称':'生姜暖贴30片','商家编码':'B','货品编号':'ITEM-B'}]).to_csv(self.single,index=False,encoding='utf-8-sig')
        self.raw = {'商品ID':'00001','skuID':'00002','商品标题':'暖贴','sku名称':'艾草暖贴30片','sku编码':'A'}
        self.write_platform()
        self.erp = v.load_erp(str(self.combo),str(self.single))
        self.records,self.index = v.build_erp_records(self.erp)
        result,meta = self.verify({})
        self.assertEqual(result.iloc[0]['核验状态'],'待人工确认')
        v.write_report(str(self.report),result,self.erp,[meta])

    def tearDown(self):
        self.temp.cleanup()

    def write_platform(self):
        pd.DataFrame([self.raw]).to_csv(self.platform,index=False,encoding='utf-8-sig')

    def verify(self,entries):
        # Retained strict resolver/import behavior; default audit has separate tests.
        with patch.object(v, 'audit_current_code', v.classify_row):
            return v.verify_platform_file(str(self.platform),self.records,self.index,confirmed_entries=entries)

    def fill(self,person='业务确认人',basis='已核对商品实物、包装与ERP记录',mark='是'):
        book=load_workbook(self.report)
        sheet=book['人工确认表']
        columns={c.value:c.column for c in sheet[1]}
        for field,value in {'确认ERP编码':'A','确认人':person,'确认依据':basis,'确认':mark}.items():
            sheet.cell(2,columns[field]).value=value
        book.save(self.report)
        book.close()

    def approve(self):
        self.fill()
        self.assertEqual(cm.import_confirmations(self.report,self.records,self.erp,self.mapping),1)
        return cm.load_mappings(self.mapping)

    def test_candidates_are_never_saved_without_explicit_confirmation(self):
        self.assertEqual(cm.import_confirmations(self.report,self.records,self.erp,self.mapping),0)
        self.assertFalse(self.mapping.exists())

    def test_confirmed_once_reused_for_same_sku(self):
        entries=self.approve()
        result,_=self.verify(entries)
        self.assertEqual(result.iloc[0]['核验状态'],'编码正确')
        self.assertEqual(result.iloc[0]['判定路径'],'人工确认映射复用')
        self.assertEqual(result.iloc[0]['人工确认人'],'业务确认人')

    def test_current_code_change_does_not_invalidate_product_confirmation(self):
        entries=self.approve()
        self.raw['sku编码']='B'
        self.write_platform()
        result,_=self.verify(entries)
        self.assertEqual(result.iloc[0]['核验状态'],'编码错误（建议修改）')
        self.assertEqual(result.iloc[0]['应改为编码'],'A')

    def test_changed_platform_spec_invalidates_confirmation(self):
        entries=self.approve()
        self.raw['sku名称']='艾草暖贴40片'
        self.write_platform()
        result,_=self.verify(entries)
        self.assertEqual(result.iloc[0]['核验状态'],'待人工确认')
        self.assertIn('重新确认',result.iloc[0]['匹配依据'])

    def test_changed_erp_record_invalidates_confirmation(self):
        entries=self.approve()
        frame=pd.read_csv(self.single,dtype=str,encoding='utf-8-sig')
        frame.loc[0,'货品名称']='艾草暖贴30片新款'
        frame.to_csv(self.single,index=False,encoding='utf-8-sig')
        self.erp=v.load_erp(str(self.combo),str(self.single))
        self.records,self.index=v.build_erp_records(self.erp)
        result,_=self.verify(entries)
        self.assertEqual(result.iloc[0]['核验状态'],'待人工确认')
        self.assertIn('ERP对应记录已变化',result.iloc[0]['匹配依据'])

    def test_missing_person_rejects_import(self):
        self.fill(person='')
        with self.assertRaisesRegex(ValueError,'确认人'):
            cm.import_confirmations(self.report,self.records,self.erp,self.mapping)
        self.assertFalse(self.mapping.exists())

    def test_missing_basis_rejects_import(self):
        self.fill(basis='')
        with self.assertRaisesRegex(ValueError,'确认依据'):
            cm.import_confirmations(self.report,self.records,self.erp,self.mapping)

    def test_modified_source_before_import_requires_new_report(self):
        self.fill()
        self.raw['商品标题']='新的商品标题'
        self.write_platform()
        with self.assertRaisesRegex(ValueError,'原始平台文件已变化'):
            cm.import_confirmations(self.report,self.records,self.erp,self.mapping)

    def test_modified_form_original_fields_are_rejected(self):
        self.fill()
        book=load_workbook(self.report)
        sheet=book['人工确认表']
        columns={c.value:c.column for c in sheet[1]}
        sheet.cell(2,columns['SKU']).value='篡改后的SKU'
        book.save(self.report);book.close()
        with self.assertRaisesRegex(ValueError,'原始商品字段被修改'):
            cm.import_confirmations(self.report,self.records,self.erp,self.mapping)

    def test_exception_status_is_kept_for_unknown_current_code(self):
        entries=self.approve()
        self.raw['sku编码']='OLD-CODE'
        self.write_platform()
        result,_=self.verify(entries)
        self.assertEqual(result.iloc[0]['核验状态'],'异常（ERP未找到）')
        self.assertEqual(result.iloc[0]['应改为编码'],'')

    def test_blank_code_with_confirmed_identity_has_separate_supplement_sheet(self):
        entries=self.approve()
        self.raw['sku编码']=''
        self.write_platform()
        result,meta=self.verify(entries)
        v.write_report(str(self.report),result,self.erp,[meta])
        missing=pd.read_excel(self.report,sheet_name='编码待补录',dtype=str,keep_default_na=False)
        self.assertEqual(len(missing),1)
        self.assertEqual(missing.iloc[0]['ERP唯一对应编码'],'A')
        self.assertEqual(missing.columns.get_loc('ERP唯一对应编码'),missing.columns.get_loc('原编码')+1)
        self.assertEqual(missing.iloc[0]['应改为编码'],'')
        self.assertEqual(len(pd.read_excel(self.report,sheet_name='待人工确认')),0)
        self.assertEqual(len(pd.read_excel(self.report,sheet_name='人工确认表')),0)

    def test_cli_import_confirms_same_workflow(self):
        self.fill()
        run=subprocess.run([sys.executable,str(Path(v.__file__)), '--combo-erp',str(self.combo),
                            '--single-erp',str(self.single),'--import-confirmations',str(self.report),
                            '--confirmed-mappings',str(self.mapping)],capture_output=True)
        self.assertEqual(run.returncode,0,run.stderr.decode('utf-8',errors='replace'))
        self.assertEqual(len(cm.load_mappings(self.mapping)),1)


if __name__=='__main__':
    unittest.main()
