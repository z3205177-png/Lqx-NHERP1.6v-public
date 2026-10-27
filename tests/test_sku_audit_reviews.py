import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import sku_audit_reviews as reviews
import verify_codes as v
from simple_report import prepare_rows, write_simple_report


class SkuAuditReviewsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        v.load_alias_config(str(ROOT / 'config/platform_aliases.json'))
        cls.records, cls.by_code = v.build_erp_records(v.load_erp(
            str(ROOT / 'data/组合装.csv'), str(ROOT / 'data/单品.csv')))
        cls.entries = reviews.load_reviews()

    def entry(self, sku):
        return next(e for e in self.entries.values() if e['context']['SKU ID'] == sku)

    def run_review(self, entry, code=None, context=None, by_code=None, entries=None):
        context = dict(context or entry['context'])
        code = entry['source_code'] if code is None else code
        index = by_code if by_code is not None else self.by_code
        pf = v.build_features(context['商品标题'], context['SKU'])
        pf['context'] = {**context, '店铺': '测试店'}
        original = v.audit_current_code(code, pf, index, self.records)
        result, proposed = reviews.apply_review(context, code, original, index,
            self.entries if entries is None else entries)
        return original, result, proposed

    def test_all_104_user_rows_follow_the_confirmed_modes(self):
        self.assertEqual(len(self.entries), 104)
        for e in self.entries.values():
            with self.subTest(sku=e['context']['SKU ID']):
                original, result, proposed = self.run_review(e)
                if e['mode'] == 'pending':
                    self.assertEqual(result['核验状态'], '待人工确认')
                    self.assertEqual(proposed['推测商家编码'], '')
                elif e['mode'] == 'conflict':
                    self.assertEqual(result['核验状态'], '异常（商品/规格不一致）')
                    self.assertIn('9英寸', result['核验说明'])
                    self.assertIn('12寸', result['核验说明'])
                    self.assertEqual(proposed['推测商家编码'], '')
                elif e['target_code'] == e['source_code']:
                    self.assertEqual(result['核验状态'], '核对通过')
                    self.assertEqual(proposed['推测商家编码'], '')
                else:
                    self.assertNotEqual(result['核验状态'], '核对通过')
                    self.assertEqual(proposed['推测商家编码'], e['target_code'])
                    self.assertEqual(result['应改为编码'], '')

    def test_all_confirmed_target_codes_pass_after_platform_is_corrected(self):
        for e in self.entries.values():
            if e['mode'] == 'confirmed':
                with self.subTest(sku=e['context']['SKU ID']):
                    _, result, _ = self.run_review(e, e['target_code'])
                    self.assertEqual(result['核验状态'], '核对通过')

    def test_four_female_child_relations_remain_pending(self):
        for sku in ['6119136462479', '6119136462475', '6119136462477', '6119136462478']:
            _, result, _ = self.run_review(self.entry(sku))
            self.assertEqual(result['核验状态'], '待人工确认')
            self.assertIn('儿童款', result['核验说明'])

    def test_only_the_reviewed_three_in_one_skus_are_exceptions(self):
        e = self.entry('100277872697338')
        original, result, _ = self.run_review(e)
        self.assertEqual(original['核验状态'], '异常（商品/规格不一致）')
        self.assertEqual(result['核验状态'], '核对通过')
        context = {**e['context'], 'SKU ID': 'unreviewed-sku'}
        original, result, proposed = self.run_review(e, context=context)
        self.assertEqual(result, original)
        self.assertIsNone(proposed)

    def test_other_product_or_platform_does_not_inherit_review(self):
        e = self.entry('6119136462472')
        for field in ('平台', '商品 ID'):
            original, result, proposed = self.run_review(e, context={**e['context'], field:'OTHER'})
            self.assertEqual(result, original)
            self.assertIsNone(proposed)

    def test_title_or_sku_change_invalidates_confirmation(self):
        e = self.entry('100277872697338')
        for field in ('商品标题', 'SKU'):
            _, result, proposed = self.run_review(e, context={**e['context'], field:e['context'][field]+'变更'})
            self.assertNotEqual(result['核验状态'], '核对通过')
            self.assertEqual(proposed['推测商家编码'], '')
            self.assertIn('重新确认', result['核验说明'])

    def test_erp_record_change_or_duplicate_invalidates_review(self):
        e = self.entry('100277872697338')
        for mutate in ('fingerprint', 'duplicate', 'missing'):
            index = copy.deepcopy(self.by_code)
            code = e['target_code']
            if mutate == 'fingerprint':index[code][0]['erp_fingerprint'] = 'changed'
            elif mutate == 'duplicate':index[code].append(index[code][0])
            else:index.pop(code)
            _, result, proposed = self.run_review(e, by_code=index)
            self.assertNotEqual(result['核验状态'], '核对通过')
            self.assertEqual(proposed['推测商家编码'], '')
            self.assertIn('重新确认', result['核验说明'])

    def test_previously_absent_original_code_appearing_requires_review(self):
        e = self.entry('5068614046742')
        index = dict(self.by_code)
        index[e['source_code']] = [index[e['target_code']][0]]
        _, _, proposed = self.run_review(e, by_code=index)
        self.assertEqual(proposed['推测商家编码'], '')

    def test_new_unreviewed_current_code_does_not_reuse_target(self):
        e = self.entry('5068614046742')
        _, result, proposed = self.run_review(e, code='OTHER')
        self.assertEqual(result['核验状态'], '异常（ERP未找到）')
        self.assertEqual(proposed['推测商家编码'], '')

    def test_literal_code_wrapper_and_blank_original_are_preserved(self):
        e = self.entry('5068614046742')
        _, result, _ = self.run_review(e, code='="'+e['target_code']+'"')
        self.assertEqual(result['核验状态'], '核对通过')
        _, result, proposed = self.run_review(e, code='')
        self.assertEqual(result['核验状态'], '编码缺失（待补录）')
        self.assertEqual(result['应改为编码'], '')
        self.assertEqual(proposed['推测商家编码'], e['target_code'])

    def test_current_quantities_and_global_business_rules_are_not_rewritten(self):
        e = self.entry('6119136462472')
        original, result, _ = self.run_review(e)
        self.assertEqual(original['数量解析'], result['数量解析'])
        self.assertFalse(v.quantity_compare((19,'对'), (20,'对'), {'发热鞋垫'}))
        self.assertEqual(v.BUSINESS_RULES['three_in_one_count'], 'total')

    def test_optional_file_and_malformed_file_are_distinct(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'reviews.json'
            self.assertEqual(reviews.load_reviews(path), {})
            path.write_text('{', encoding='utf-8')
            with self.assertRaises(json.JSONDecodeError):reviews.load_reviews(path)

    def test_full_pipeline_and_simple_workbook_expose_pending_reason(self):
        e = self.entry('6119136462472')
        c = e['context']
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder)/'快手.csv', Path(folder)/'result.xlsx'
            pd.DataFrame([{'商品Id':c['商品 ID'],'skuId':c['SKU ID'],
                '商品标题':c['商品标题'],'销售属性':c['SKU'],'商家编码.1':e['source_code']}]).to_csv(source,index=False,encoding='utf-8-sig')
            rows, _ = v.verify_platform_file(str(source), self.records, self.by_code, platform_override='天猫')
            self.assertEqual(rows.iloc[0]['核验状态'], '待人工确认')
            self.assertEqual(rows.iloc[0]['原编码'], e['source_code'])
            write_simple_report(output, rows)
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ['需处理','全部结果'])
            self.assertEqual(book['全部结果'].max_column, 12)
            self.assertEqual(book['全部结果']['G5'].value, '待确认')
            self.assertIn('未认定', book['全部结果']['I5'].value)
            self.assertIn(book['全部结果']['H5'].value, (None, ''))
            book.close()


if __name__ == '__main__':
    unittest.main()
