import sys
import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import verify_codes as v
from test_matching import erp_record


class CandidateSuggestionTests(unittest.TestCase):
    def propose(self, sku, names, title='暖贴', status='异常（ERP未找到）'):
        records = [erp_record(code, name) for code, name in names]
        for i, record in enumerate(records):
            record.update(erp_fingerprint=str(i), source_row=i+2)
        by_code = {}
        for record in records:
            by_code.setdefault(record['code'], []).append(record)
        pf = v.build_features(title, sku)
        pf['context'] = {'平台':'测试平台','商品 ID':'1','商品标题':title,'SKU':sku}
        result = {'核验状态':status,'应改为编码':''}
        original = result.copy()
        proposed = v.suggest_code_candidates(pf, result, records, by_code)
        self.assertEqual(result, original)
        return proposed

    def test_absent_code_has_unique_real_erp_candidate(self):
        result = self.propose('艾草暖贴60片',[('000001','艾草暖贴60片'),('000002','艾草暖贴100片')])
        self.assertEqual(result['推测商家编码'],'000001')
        self.assertEqual(result['推测候选编码数'],1)

    def test_wrong_quantity_candidate_does_not_repeat_wrong_code(self):
        result = self.propose('艾草暖贴60片',[('RIGHT','艾草暖贴60片'),('WRONG','艾草暖贴100片')],status='异常（商品/规格不一致）')
        self.assertEqual(result['推测商家编码'],'RIGHT')

    def test_multiple_equal_candidates_leave_primary_blank(self):
        result = self.propose('艾草暖贴60片',[('A','暖火艾草暖贴60片'),('B','艾草暖贴60片')])
        self.assertEqual(result['推测商家编码'],'')
        self.assertEqual(result['其他候选编码'],'A / B')
        self.assertIn('暖火艾草暖贴60片',result['其他候选ERP记录'])

    def test_passed_rows_are_not_given_replacement_codes(self):
        result = self.propose('暖贴60片',[('A','暖贴60片')],status='核对通过')
        self.assertEqual(result['推测商家编码'],'')
        self.assertEqual(result['推测候选编码数'],0)

    def test_no_matching_spec_is_blank(self):
        result = self.propose('艾草暖贴60片',[('A','艾草暖贴100片')])
        self.assertEqual(result['推测候选编码数'],0)
        self.assertEqual(result['推测商家编码'],'')

    def test_three_in_one_does_not_use_single_formula(self):
        result = self.propose('三合一暖贴10片',[('A','艾草暖贴10片')])
        self.assertEqual(result['推测候选编码数'],0)

    def test_glove_color_material_size_and_count(self):
        result = self.propose('蓝色PVC手套M码50只',[
            ('A','蓝色PVC手套M码50只'),('B','透明PVC手套M码50只'),
            ('C','蓝色丁腈手套M码50只'),('D','蓝色PVC手套L码50只'),('E','蓝色PVC手套M码100只')],title='手套')
        self.assertEqual(result['推测商家编码'],'A')

    def test_single_glove_does_not_use_mixed_material_bundle(self):
        result=self.propose('尺码:S;颜色分类:100只蓝色PVC;', [('A','PVC手套S码100只+TPE蓝色手套S码100只')],title='PVC手套',status='编码缺失（待补录）')
        self.assertEqual(result['推测候选编码数'],0)

    def test_knee_primary_kind_can_come_from_title(self):
        result=self.propose('咖色一对【搭配15片艾草暖贴】,L码', [('A','护膝咖色L+艾草暖贴15片')],title='保暖护膝')
        self.assertEqual(result['推测商家编码'],'A')

    def test_knee_color_size_and_paired_quantity(self):
        result = self.propose('咖色护膝L搭配艾草暖贴15片',[
            ('A','护膝咖色L+艾草暖贴15片'),('B','护膝灰色L+艾草暖贴15片'),
            ('C','护膝咖色XL+艾草暖贴15片'),('D','护膝咖色L+艾草暖贴5片')],title='护膝')
        self.assertEqual(result['推测商家编码'],'A')

    def test_unknown_939_letter_size_is_not_inferred(self):
        result = self.propose('尺码:S;颜色分类:高筒透明色PVC鞋套;',[
            ('A','防雨鞋套939透明 34/35码'),('B','防雨鞋套919透明 30/31码')],title='鞋套',status='编码缺失（待补录）')
        self.assertEqual(result['推测候选编码数'],0)

    def test_duplicate_erp_codes_are_not_guessed(self):
        result = self.propose('艾草暖贴60片',[('A','艾草暖贴60片'),('A','艾草暖贴100片')])
        self.assertEqual(result['推测候选编码数'],0)

    def test_name_similarity_alone_is_not_enough(self):
        result = self.propose('优惠经济装',[('A','暖贴60片')])
        self.assertEqual(result['推测候选编码数'],0)

    def test_pair_conversion_candidate(self):
        result = self.propose('足贴60片',[('A','足贴艾草30对'),('B','足贴艾草60对')],title='足贴')
        self.assertEqual(result['推测商家编码'],'A')

    def test_child_specification_in_erp_name_is_supporting_gender(self):
        result=self.propose('儿童款10片鞋垫',[('A','鞋垫儿童5对'),('B','鞋垫男码5对')],title='鞋垫')
        self.assertEqual(result['推测商家编码'],'A')

    def test_known_supplement_target_is_directly_reused(self):
        record=erp_record('A','艾草暖贴60片')
        result=v.suggest_code_candidates(v.build_features('暖贴','经济装'),
            {'核验状态':'编码缺失（待补录）','ERP唯一对应编码':'A'},[record],{'A':[record]})
        self.assertEqual(result['推测商家编码'],'A')

    def test_saved_workbook_has_review_columns_and_preserves_original_status(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            combo,single,platform,report=[root/name for name in ('combo.csv','single.csv','shop.csv','result.xlsx')]
            pd.DataFrame(columns=['组合装名称','商家编码']).to_csv(combo,index=False,encoding='utf-8-sig')
            pd.DataFrame([{'货品名称':'艾草暖贴60片','商家编码':'000123'}]).to_csv(single,index=False,encoding='utf-8-sig')
            pd.DataFrame([{'商品ID':'001','skuID':'002','商品标题':'暖贴','sku名称':'艾草暖贴60片','sku编码':'OLD'}]).to_csv(platform,index=False,encoding='utf-8-sig')
            erp=v.load_erp(str(combo),str(single))
            records,index=v.build_erp_records(erp)
            rows,meta=v.verify_platform_file(str(platform),records,index,confirmed_entries={})
            v.write_report(str(report),rows,erp,[meta])
            book=load_workbook(report)
            sheet=book['核验结果']
            headers={cell.value:cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2,headers['推测商家编码']).value,'000123')
            self.assertEqual(sheet.cell(2,headers['推测商家编码']).data_type,'s')
            self.assertEqual(sheet.cell(2,headers['原编码']).value,'OLD')
            self.assertEqual(sheet.cell(2,headers['核验状态']).value,'异常（ERP未找到）')
            self.assertIsNone(sheet.cell(2,headers['应改为编码']).value)
            self.assertGreater(headers['推测商家编码'],headers['原编码'])
            self.assertEqual(book['异常及补录推测'].max_row,2)
            detail=book['推测候选明细']
            details={cell.value:cell.column for cell in detail[1]}
            self.assertEqual(detail.cell(2,details['候选ERP编码']).value,'000123')
            book.close()

    def test_confirmed_style_binding_is_scoped_and_fingerprinted(self):
        records = [erp_record('A','KT暖贴60片'),erp_record('B','卡通暖贴60片')]
        records[0]['erp_fingerprint']='fp-a'
        records[1]['erp_fingerprint']='fp-b'
        pf=v.build_features('指定品牌暖贴','卡通暖贴60片')
        pf['context']={'平台':'测试平台','商品 ID':'1','商品标题':'指定品牌暖贴'}
        rule={'平台':'测试平台','商品 ID':'1','商品标题包含':'指定品牌','平台SKU':'卡通暖贴60片','ERP编码':'A','ERP记录指纹':'fp-a'}
        index={r['code']:[r] for r in records}
        with patch.dict(v.BUSINESS_RULES,{'confirmed_style_bindings':[rule]}):
            result=v.suggest_code_candidates(pf,{'核验状态':'异常（ERP未找到）'},records,index)
            self.assertEqual(result['推测商家编码'],'A')
            pf['context']['商品 ID']='2'
            result=v.suggest_code_candidates(pf,{'核验状态':'异常（ERP未找到）'},records,index)
            self.assertEqual(result['推测商家编码'],'')
            self.assertEqual(result['推测候选编码数'],2)
            pf['context']['商品 ID']='1'
            records[0]['erp_fingerprint']='changed'
            result=v.suggest_code_candidates(pf,{'核验状态':'异常（ERP未找到）'},records,index)
            self.assertEqual(result['推测商家编码'],'')


if __name__=='__main__':
    unittest.main()
