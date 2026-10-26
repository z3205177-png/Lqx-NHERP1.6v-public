import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import verify_codes as v
import erp_library as library
from simple_report import write_simple_report, prepare_rows, HEADERS
from test_matching import erp_record


class RegisteredLibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        (self.root/'config').mkdir()
        (self.root/'data').mkdir()
        self.combo,self.single=[self.root/'data'/name for name in ('combo.csv','single.csv')]
        self.combo.write_text('商家编码,组合装名称\nA,暖贴60片\n',encoding='utf-8')
        self.single.write_text('商家编码,货品名称\nB,暖贴30片\n',encoding='utf-8')
        self.manifest=self.root/'config/library.json'
        self.manifest.write_text(json.dumps({'version':'test-v1','erp':{key:{'path':str(path.relative_to(self.root)),
            'sha256':hashlib.sha256(path.read_bytes()).hexdigest()} for key,path in [('combo',self.combo),('single',self.single)]}}),encoding='utf-8')

    def tearDown(self):
        self.temp.cleanup()

    def test_default_uses_registered_pair(self):
        combo,single,meta=library.resolve_erp_inputs(manifest_path=self.manifest)
        self.assertTrue(Path(combo).is_absolute())
        self.assertTrue(Path(single).is_absolute())
        self.assertTrue(Path(combo).samefile(self.combo))
        self.assertTrue(Path(single).samefile(self.single))
        self.assertEqual(meta['version'],'test-v1')

    def test_library_is_relative_to_project_not_current_directory(self):
        self.assertTrue(Path(library.resolve_erp_inputs(manifest_path=self.manifest)[0]).is_absolute())

    def test_modified_library_requires_explicit_update(self):
        before=self.manifest.read_bytes()
        self.combo.write_text('changed',encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'登记版本不一致'):
            library.resolve_erp_inputs(manifest_path=self.manifest)
        self.assertEqual(before,self.manifest.read_bytes())

    def test_missing_library_file_is_reported(self):
        self.combo.unlink()
        with self.assertRaisesRegex(FileNotFoundError,'固定ERP标准库文件缺失'):
            library.resolve_erp_inputs(manifest_path=self.manifest)

    def test_one_override_is_rejected(self):
        with self.assertRaises(ValueError):
            library.resolve_erp_inputs(combo=self.combo,manifest_path=self.manifest)

    def test_explicit_pair_does_not_update_registered_library(self):
        before=self.manifest.read_bytes()
        _,_,meta=library.resolve_erp_inputs(self.single,self.combo,manifest_path=self.manifest)
        self.assertEqual(meta['mode'],'one_run_override')
        self.assertEqual(before,self.manifest.read_bytes())


class SimpleDeliveryTests(unittest.TestCase):
    def row(self,sku='暖贴60片',state='异常（ERP未找到）',item='001'):
        return {'平台':'快手','商品 ID':item,'SKU ID':'00002','商品标题':'暖贴','SKU':sku,
            '原编码':'000OLD','核验状态':state,'推测商家编码':'','候选明细JSON':'[]','冲突项目':''}

    def test_candidate_cells_show_code_and_original_erp_name(self):
        row=self.row()
        row['候选明细JSON']=json.dumps([{'code':'000123','name':'暖火艾草暖贴60片'},{'code':'000124','name':'艾草暖贴60片'}])
        pending,_=prepare_rows(pd.DataFrame([row]))
        self.assertEqual(pending[0]['候选1'],'000123\n暖火艾草暖贴60片')
        self.assertEqual(pending[0]['推测编码'],'')

    def test_only_top_three_are_shown_with_total_count(self):
        row=self.row()
        row['候选明细JSON']=json.dumps([{'code':str(i),'name':'暖贴60片'} for i in range(5)])
        pending,_=prepare_rows(pd.DataFrame([row]))
        self.assertIn('5个候选',pending[0]['问题说明'])
        self.assertEqual(pending[0]['候选3'],'2\n暖贴60片')

    def test_conflicts_first_same_specifications_grouped_and_no_rows_deleted(self):
        rows=[self.row('60片','异常（ERP未找到）','1'),self.row('30片','异常（商品/规格不一致）','2'),self.row('60片','异常（ERP未找到）','3'),self.row('30片','核对通过','4')]
        pending,all_rows=prepare_rows(pd.DataFrame(rows))
        self.assertEqual([r['商品ID'] for r in pending],['2','1','3'])
        self.assertEqual([r['商品ID'] for r in all_rows],['1','2','3','4'])

    def test_native_workbook_has_two_sheets_and_text_codes(self):
        row=self.row()
        row.update(商品标题='=1+1',推测商家编码='000123')
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'simple.xlsx'
            write_simple_report(path,pd.DataFrame([row]))
            book=load_workbook(path,data_only=False)
            self.assertEqual(book.sheetnames,['需处理','全部结果'])
            self.assertEqual(book.active.title,'需处理')
            sheet=book['全部结果']
            self.assertEqual([c.value for c in sheet[4]],HEADERS)
            self.assertEqual(sheet['B5'].value,'001')
            self.assertEqual(sheet['H5'].value,'000123')
            self.assertEqual(sheet['D5'].value,'=1+1')
            self.assertEqual(sheet['D5'].data_type,'s')
            book.close()

    def test_all_passed_still_has_an_empty_pending_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'simple.xlsx'
            write_simple_report(path,pd.DataFrame([self.row(state='核对通过')]))
            book=load_workbook(path)
            self.assertEqual(book['需处理'].max_row,4)
            self.assertEqual(book['全部结果'].max_row,5)
            book.close()

    def test_cli_needs_only_platform_file(self):
        with tempfile.TemporaryDirectory() as folder:
            platform=Path(folder)/'快手.csv'
            output=Path(folder)/'result.xlsx'
            pd.DataFrame([{'商品ID':'001','skuID':'002','商品标题':'暖贴','sku名称':'描述不全','sku编码':'NHCP00801'}]).to_csv(platform,index=False,encoding='utf-8-sig')
            before=library.DEFAULT_MANIFEST.read_bytes()
            completed=subprocess.run([sys.executable,str(Path(v.__file__)),'--platform-file',str(platform),'--output',str(output)],capture_output=True)
            self.assertEqual(completed.returncode,0,completed.stderr.decode('utf-8',errors='replace'))
            book=load_workbook(output)
            self.assertEqual(book.sheetnames,['需处理','全部结果'])
            self.assertEqual(book['全部结果']['G5'].value,'通过')
            self.assertTrue((output.parent/'核验依据/result.json').is_file())
            self.assertEqual(before,library.DEFAULT_MANIFEST.read_bytes())
            book.close()


class CandidateOrderingTests(unittest.TestCase):
    def test_explicit_evidence_changes_order_but_does_not_choose_a_single_code(self):
        with patch.dict(v.BUSINESS_RULES,{'confirmed_style_bindings':[],'confirmed_code_sets':[]}):
            pf=v.build_features('暖贴','艾草暖贴60片发热12小时')
            records=[erp_record('A','艾草暖贴60片'),erp_record('Z','艾草暖贴60片发热12小时')]
            result=v.suggest_code_candidates(pf,{'核验状态':'异常（ERP未找到）'},records,{r['code']:[r] for r in records})
        self.assertEqual(result['推测商家编码'],'')
        self.assertEqual(result['其他候选编码'],'Z / A')
        evidence=json.loads(result['候选明细JSON'])
        self.assertGreater(len(evidence[0]['evidence']),len(evidence[1]['evidence']))
