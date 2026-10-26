"""Regression coverage for Windows checkout and pandas report preparation."""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import erp_library


class CheckoutCompatibilityTests(unittest.TestCase):
    @unittest.skipUnless(shutil.which('git'), 'Git is only needed for checkout verification')
    def test_erp_bytes_survive_checkout_with_either_autocrlf_setting(self):
        manifest = json.loads((ROOT / 'config/erp_library.json').read_text(encoding='utf-8'))
        sources = ['config/erp_library.json', *[v['path'] for v in manifest['erp'].values()]]
        with tempfile.TemporaryDirectory() as folder:
            for setting in ('true', 'false'):
                with self.subTest(autocrlf=setting):
                    staging = Path(folder) / ('staging-' + setting)
                    checkout = Path(folder) / ('checkout-' + setting)
                    staging.mkdir()
                    checkout.mkdir()
                    for relative in [*sources, '.gitattributes']:
                        source = ROOT / relative
                        if source.is_file():
                            target = staging / relative
                            target.parent.mkdir(parents=True, exist_ok=True)
                            shutil.copyfile(source, target)
                    commands = [
                        ['git', 'init', '--quiet'],
                        ['git', 'config', 'core.autocrlf', setting],
                        ['git', 'add', '--all'],
                        ['git', 'checkout-index', '--all', '--prefix=' + checkout.as_posix() + '/'],
                    ]
                    for command in commands:
                        subprocess.run(command, cwd=staging, check=True, capture_output=True)
                    combo, single, info = erp_library.resolve_erp_inputs(
                        manifest_path=checkout / 'config/erp_library.json')
                    self.assertEqual(info['version'], manifest['version'])
                    for actual, key in ((combo, 'combo'), (single, 'single')):
                        self.assertEqual(Path(actual).read_bytes(),
                                         (ROOT / manifest['erp'][key]['path']).read_bytes())


class PandasCompatibilityTests(unittest.TestCase):
    def test_cli_mixed_status_rows_preserve_values_without_dtype_warnings(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            combo, single, platform, output = [base / name for name in
                ('combo.csv', 'single.csv', 'shop.csv', 'result.xlsx')]
            pd.DataFrame(columns=['组合装名称', '商家编码']).to_csv(combo, index=False, encoding='utf-8-sig')
            pd.DataFrame([{'货品名称': '暖贴30片', '商家编码': '000123'}]).to_csv(
                single, index=False, encoding='utf-8-sig')
            rows = [
                {'商品ID': '00001', 'skuID': '00011', '商品标题': '=1+1', 'sku名称': '暖贴30片', 'sku编码': '000123'},
                {'商品ID': '00002', 'skuID': '00012', '商品标题': '暖贴', 'sku名称': '暖贴40片', 'sku编码': '000123'},
                {'商品ID': '00003', 'skuID': '00013', '商品标题': '暖贴', 'sku名称': '暖贴30片', 'sku编码': 'MISSING'},
                {'商品ID': '00004', 'skuID': '00014', '商品标题': '暖贴', 'sku名称': '暖贴30片', 'sku编码': ''},
            ]
            pd.DataFrame(rows).to_csv(platform, index=False, encoding='utf-8-sig')
            env = os.environ.copy()
            env['PYTHONUTF8'] = '1'
            command = [sys.executable, '-W', 'error::DeprecationWarning',
                       str(ROOT / 'scripts/verify_codes.py'), '--combo-erp', str(combo),
                       '--single-erp', str(single), '--platform-file', str(platform),
                       '--output', str(output)]
            completed = subprocess.run(command, capture_output=True, encoding='utf-8', env=env)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertNotIn('Pandas4Warning', completed.stderr)
            audit = json.loads((base / '核验依据/result.json').read_text(encoding='utf-8'))
            results = audit['results']
            self.assertEqual([r['核验状态'] for r in results],
                ['核对通过', '异常（商品/规格不一致）', '异常（ERP未找到）', '编码缺失（待补录）'])
            self.assertEqual([r['原编码'] for r in results], [r['sku编码'] for r in rows])
            self.assertEqual(results[0]['当前编码ERP记录数'], 1)
            self.assertEqual(results[3]['当前编码ERP记录数'], '')
            book = load_workbook(output)
            try:
                self.assertEqual(book.sheetnames, ['需处理', '全部结果'])
                self.assertEqual(book['全部结果'].max_column, 12)
                self.assertEqual(book['全部结果']['B5'].value, '00001')
                self.assertEqual(book['全部结果']['F5'].value, '000123')
                self.assertEqual(book['全部结果']['D5'].value, '=1+1')
                self.assertEqual(book['全部结果']['D5'].data_type, 's')
            finally:
                book.close()
