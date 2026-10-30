"""统一回归入口：多平台批量核验 + 基线对比 + 覆盖情况报告。

设计约束：
1. 样本路径属于本机私有资料，放在外部配置里，不进仓库。
2. **缺样本必须显式报告**。只跑了 4 个平台就写「全部通过」是不能接受的，
   所以输出末尾固定打印覆盖情况，缺哪个、为什么缺，都要说清楚。
3. 基线可选。没有基线时只输出本次结果，不宣称「无变化」。

用法：

    python scripts/regress_all.py --config 回归配置.json --output outputs/回归_20261030
    python scripts/regress_all.py --config 回归配置.json --baseline outputs/回归_p7

配置格式见 config/regression_samples.example.json。
"""
from __future__ import annotations

import argparse
import json
import hashlib
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
STATUS_COLUMNS = ['平台', '商品ID', 'SKU_ID', '名称', '规格', '原编码',
                  '状态', '推测', '问题', 'c1', 'c2', 'c3']
COMPARE_FIELDS = ('平台', '名称', '规格', '状态', '原编码', '推测', 'c1', 'c2', 'c3')
STATUS_KEYS = ('通过', '不一致', '未找到', '空码', '待确认')

# 交付 xlsx 的短标签和内部 JSON 的长状态都要识别。
STATUS_ALIASES = {
    '核对通过': '通过', '通过': '通过',
    '异常（商品/规格不一致）': '不一致', '规格不一致': '不一致', '不一致': '不一致',
    '异常（ERP未找到）': '未找到', 'ERP未找到': '未找到', '未找到': '未找到',
    '编码缺失（待补录）': '空码', '编码为空': '空码', '空码': '空码',
    '待人工确认': '待确认', '待确认': '待确认',
}


def normalize_status(value):
    """把长状态、短状态统一成内部短名；无法识别时返回原文并留待上报。"""
    text = str(value or '').strip()
    return STATUS_ALIASES.get(text, text)


# ---------------------------------------------------------------- 运行

def run_platform(item, output_dir, python_exe):
    """跑一个平台的核验，返回输出路径或失败原因。"""
    files = [Path(f) for f in item.get('files', [])]
    missing = [str(f) for f in files if not f.is_file()]
    if not files:
        return None, '配置里没有样本文件'
    if missing:
        return None, '文件不存在：' + '、'.join(missing)

    hook = item.get('preprocess')
    if hook:
        return None, '不支持修改待核编码的预处理；请使用已确认字段映射或注明来源的历史模板'

    output = output_dir / f"{item['name']}.xlsx"
    if output.resolve() in [f.resolve() for f in files]:
        return None, '输出路径与源文件相同，拒绝覆盖原始输入'
    cmd = [python_exe, str(PROJECT_ROOT / 'scripts' / 'verify_codes.py')]
    for f in files:
        cmd += ['--platform-file', str(f), '--platform', item['platform'], '--store', item['store']]
        if item.get('sheet') is not None:
            cmd += ['--sheet', str(item['sheet'])]
    cmd += ['--output', str(output)]

    try:
        result = subprocess.run(cmd, capture_output=True, text=True,
                                encoding='utf-8', errors='replace', timeout=1800)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return None, '执行失败：' + str(exc)
    if result.returncode != 0 or not output.exists():
        tail = (result.stderr or result.stdout or '').strip().splitlines()[-3:]
        return None, '核验失败：' + ' / '.join(tail)
    return output, ''


def load_status(path):
    from openpyxl import load_workbook
    from simple_report import HEADERS, prepare_rows
    path = Path(path)
    book = load_workbook(path, read_only=True, data_only=False)
    records = []
    try:
        sheet = book['全部结果']
        if [c.value for c in sheet[4]][:12] != HEADERS:
            raise ValueError('全部结果表头不是预期的12列')
        for row in sheet.iter_rows(min_row=5, max_col=12, values_only=True):
            data = dict(zip(STATUS_COLUMNS, row))
            if not any(value is not None and str(value).strip() for value in data.values()):
                continue
            data['状态'] = normalize_status(data['状态'])
            data['_full_evidence'] = False
            records.append(data)
    finally:
        book.close()
    evidence = path.parent / '核验依据' / (path.stem + '.json')
    if evidence.is_file():
        raw = json.loads(evidence.read_text(encoding='utf-8'))['results']
        if len(raw) != len(records):
            raise ValueError('报告与核验依据行数不同')
        import pandas as pd
        _, rendered = prepare_rows(pd.DataFrame(raw))
        for data, original, visible in zip(records, raw, rendered):
            if (str(data['商品ID'] or ''), str(data['SKU_ID'] or '')) != (str(original['商品 ID'] or ''), str(original['SKU ID'] or '')):
                raise ValueError('报告与核验依据ID或顺序不同')
            for field, header in zip(STATUS_COLUMNS, HEADERS):
                expected = normalize_status(visible[header]) if field == '状态' else visible[header]
                if str(data.get(field) or '') != str(expected or ''):
                    raise ValueError('报告与核验依据字段不同：' + header)
            candidates = json.loads(original.get('候选明细JSON') or '[]')
            data['_candidates'] = sorted((c.get('code', ''), c.get('name', ''), c.get('source', ''), str(c.get('source_row', ''))) for c in candidates)
            data['_full_evidence'] = True
    return records


def compare_with_baseline(current, baseline, full=False):
    """按 (平台, 商品ID, SKU_ID) 分组比对，保留重复行。"""
    def group(rows):
        grouped = {}
        for row in rows:
            grouped.setdefault((str(row['平台']), str(row['商品ID']), str(row['SKU_ID'])), []).append(row)
        return grouped

    old_map, new_map = group(baseline), group(current)
    changes = []
    for key in sorted(set(old_map) | set(new_map)):
        old_rows, new_rows = old_map.get(key, []), new_map.get(key, [])
        sample = (old_rows or new_rows)[0]
        if len(old_rows) != len(new_rows):
            changes.append({'code': sample['原编码'], 'field': '行数',
                            'platform': key[0], 'product_id': key[1], 'sku_id': key[2],
                            'old': f'{len(old_rows)} 行', 'new': f'{len(new_rows)} 行'})
            continue
        for occurrence, (old_row, new_row) in enumerate(zip(old_rows, new_rows), 1):
            for field in (*COMPARE_FIELDS, *(('_candidates',) if full else ())):
                before = str(old_row.get(field, '') or '')
                after = str(new_row.get(field, '') or '')
                if before != after:
                    changes.append({'code': old_row['原编码'], 'field': field,
                                    'platform': key[0], 'product_id': key[1], 'sku_id': key[2],
                                    'occurrence': occurrence, 'old': before, 'new': after})
    identities = lambda rows: [(str(r['平台']), str(r['商品ID']), str(r['SKU_ID'])) for r in rows]
    old_order, new_order = identities(baseline), identities(current)
    if Counter(old_order) == Counter(new_order) and old_order != new_order:
        changes.append({'code': '', 'field': '行序', 'old': old_order, 'new': new_order})
    return changes


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description='多平台批量回归（缺样本会显式报告）')
    parser.add_argument('--config', required=True, help='回归配置 JSON')
    parser.add_argument('--output', help='输出目录，默认为 outputs/回归_<时间戳>')
    parser.add_argument('--baseline', help='基线目录（含各平台 xlsx）')
    parser.add_argument('--python', default=sys.executable, help='解释器路径')
    parser.add_argument('--only', action='append', default=[], help='只跑指定样本名，可重复')
    parser.add_argument('--fail-on-change', action='store_true', help='比较有变化时退出码1')
    parser.add_argument('--comparison-scope', choices=['full', 'report'], default='full',
                        help='full要求完整候选依据；report仅比报告可见字段')
    args = parser.parse_args(argv)

    config = json.loads(Path(args.config).read_text(encoding='utf-8'))
    platforms = config.get('platforms', [])
    configured_samples = len(platforms)
    names = [p['name'] for p in platforms]
    if len(names) != len(set(names)):
        parser.error('样本name重复')
    if set(args.only) - set(names):
        parser.error('--only含不存在的样本名')
    for item in platforms:
        if not item['name'] or item['name'] in {'.', '..'} or any(c in item['name'] for c in '/\\:*?"<>|'):
            parser.error('样本name不能含路径或非法文件名字符')
        if not item.get('platform') or not item.get('store'):
            parser.error('每个样本必须填写platform和store')
    if args.only:
        platforms = [p for p in platforms if p['name'] in args.only]
    if not platforms:
        raise SystemExit('配置里没有可运行的平台')

    stamp = datetime.now().strftime('%Y%m%d_%H%M')
    output_dir = Path(args.output or f'outputs/回归_{stamp}').resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    baseline_dir = Path(args.baseline).resolve() if args.baseline else None
    if baseline_dir == output_dir:
        parser.error('输出目录不能与基线目录相同，避免覆盖冻结结果')

    print('=' * 78)
    print(f'多平台回归  输出：{output_dir}')
    print(f'基线：{baseline_dir if baseline_dir else "未提供（只输出本次结果，不判断变化）"}')
    print('=' * 78)

    results, failures, comparison_failures = [], [], []
    for item in platforms:
        item = dict(item)
        item['files'] = [str((Path(args.config).resolve().parent / p).resolve()) for p in item.get('files', [])]
        name = item['name']
        print(f'\n[{name}] 运行中…')
        try:
            fingerprints = {p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
                            for p in item['files'] if Path(p).is_file()}
            path, error = run_platform(item, output_dir, args.python)
        except Exception as exc:
            failures.append((name, '执行失败：' + str(exc)))
            continue
        if error:
            print(f'  跳过：{error}')
            failures.append((name, error))
            continue
        try:
            rows = load_status(path)
            if not rows:
                raise ValueError('样本输出为空，不能作为完成覆盖')
            if fingerprints != {p: hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in fingerprints}:
                raise ValueError('运行过程中源文件发生变化')
        except Exception as exc:
            print('  失败：报告读取失败：' + str(exc))
            failures.append((name, '报告读取失败：' + str(exc)))
            continue
        counts = Counter(row['状态'] for row in rows)
        changes = []
        compared = False
        if baseline_dir:
            base_path = baseline_dir / f'{name}.xlsx'
            if base_path.is_file():
                try:
                    prior = load_status(base_path)
                    if any(r['状态'] not in STATUS_KEYS for r in prior):
                        raise ValueError('基线含未知状态')
                    if args.comparison_scope == 'full' and not all(r.get('_full_evidence') for r in [*rows, *prior]):
                        raise ValueError('缺少完整候选核验依据；仅比较可见列需显式选择 --comparison-scope report')
                    changes = compare_with_baseline(rows, prior, args.comparison_scope == 'full')
                    compared = True
                except Exception as exc:
                    comparison_failures.append((name, str(exc)))
            else:
                print(f'  提示：基线缺少 {name}.xlsx，本次不判断变化')
                comparison_failures.append((name, '缺少基线：' + str(base_path)))
        results.append({'name': name, 'rows': len(rows), 'counts': counts,
                        'changes': changes, 'compared': compared,
                        'platform': item['platform'], 'store': item['store'],
                        'source_kind': item.get('source_kind', '未注明'),
                        'source_sha256': fingerprints})
        suffix = f'变化 {len(changes)} 处' if compared else '未对比基线'
        print(f'  完成：{len(rows)} 行，{suffix}')

    # ---------------- 汇总
    print('\n' + '=' * 78)
    print('汇总')
    print('=' * 78)
    header = f"{'平台':<10}{'行数':>6}{'通过':>7}{'不一致':>7}{'未找到':>7}{'空码':>6}{'待确认':>7}{'变化':>6}"
    print(header)
    print('-' * len(header))
    total = Counter()
    total_rows = 0
    total_changes = 0
    compared_count = 0
    for item in results:
        c = item['counts']
        total_rows += item['rows']
        total_changes += len(item['changes'])
        compared_count += 1 if item['compared'] else 0
        for key, value in c.items():
            total[key] += value
        # 没对比过就写 '-'：写 0 会被读成「确认无变化」，那是误导
        change_cell = str(len(item['changes'])) if item['compared'] else '-'
        print(f"{item['name']:<10}{item['rows']:>6}{c.get('通过', 0):>7}{c.get('不一致', 0):>7}"
              f"{c.get('未找到', 0):>7}{c.get('空码', 0):>6}{c.get('待确认', 0):>7}{change_cell:>6}")
    print('-' * len(header))
    all_compared = bool(results) and compared_count == len(results)
    total_change_cell = str(total_changes) if all_compared else f'{total_changes}*' if compared_count else '-'
    print(f"{'合计':<10}{total_rows:>6}{total.get('通过', 0):>7}{total.get('不一致', 0):>7}"
          f"{total.get('未找到', 0):>7}{total.get('空码', 0):>6}{total.get('待确认', 0):>7}"
          f"{total_change_cell:>6}")
    if baseline_dir and not all_compared:
        print(f"  * 合计变化数只统计了已对比的 {compared_count}/{len(results)} 个样本")

    unknown = {key: value for key, value in total.items() if key not in STATUS_KEYS}
    if unknown:
        print(f'\n⚠ 出现未识别的状态，下面的分类统计不完整：{unknown}')
        print('  请把它们补进 STATUS_ALIASES，否则缺口会被静默忽略。')

    # ---------------- 变化明细
    changed = [item for item in results if item['changes']]
    if changed:
        print('\n变化明细（按平台）')
        for item in changed:
            print(f"\n  [{item['name']}] {len(item['changes'])} 处")
            for change in item['changes'][:12]:
                print(f"    {change['code']}  [{change['field']}] {change['old']} → {change['new']}")
            if len(item['changes']) > 12:
                print(f"    … 共 {len(item['changes'])} 处，仅列前 12")

    # ---------------- 覆盖情况（核心：不许隐瞒缺样本）
    print('\n' + '=' * 78)
    print('覆盖情况')
    print('=' * 78)
    configured = len(platforms)
    done = len(results)
    print(f'配置样本 {configured} 个，实际完成 {done} 个，缺失 {len(failures)} 个')
    for name, error in failures:
        print(f'  缺失：{name} —— {error}')
    if failures:
        print(f'\n结论：本次仅覆盖 {done}/{configured} 个样本，'
              f'**不能表述为全量通过**。缺失样本需补齐后重跑。')
    else:
        print(f'\n结论：覆盖 {configured}/{configured} 个样本，共 {total_rows} 行。')
    if not baseline_dir:
        print('注意：未提供基线，本次没有做变化比对。')
    if args.only:
        print(f'本次为选择性回归：选择 {configured}/{configured_samples} 个样本，不代表全部配置范围。')
    for name, error in comparison_failures:
        print(f'未完成比较：{name} —— {error}')
    incomplete = bool(failures or comparison_failures or unknown)
    if incomplete:
        print('不能表述为全量通过：样本、基线覆盖或状态统计不完整。')
    comparison_complete = bool(baseline_dir) and not incomplete and all_compared
    exit_code = 2 if incomplete else 1 if args.fail_on_change and total_changes else 0

    report = {
        'generated_at': datetime.now().isoformat(timespec='seconds'),
        'output_dir': str(output_dir),
        'baseline': str(baseline_dir) if baseline_dir else None,
        'configured': configured, 'completed': done,
        'failures': [{'name': n, 'reason': r} for n, r in failures],
        'platforms': [{'name': i['name'], 'rows': i['rows'],
                       'counts': dict(i['counts']),
                       'baseline_compared': i['compared'],
                       'platform': i['platform'], 'store': i['store'],
                       'source_kind': i['source_kind'], 'source_sha256': i['source_sha256'],
                       'changes': i['changes']} for i in results],
        'total_rows': total_rows, 'total_changes': total_changes if comparison_complete else None,
        'compared_changes': total_changes, 'counts': dict(total),
        'coverage_complete': not failures,
        'comparison_complete': comparison_complete,
        'status_valid': not unknown,
        'unknown_statuses': unknown,
        'comparison_failures': [{'name': n, 'reason': r} for n, r in comparison_failures],
        'configured_samples': configured_samples, 'selected_samples': configured,
        'selection_is_full': configured == configured_samples,
        'comparison_scope': args.comparison_scope,
        'exit_code': exit_code,
    }
    report_path = output_dir / '回归报告.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'\n报告：{report_path}')
    return exit_code


if __name__ == '__main__':
    sys.exit(main())
