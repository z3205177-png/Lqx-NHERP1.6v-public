#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Audit current merchant-code registration and explicit product/spec conflicts.

Usage:
  python scripts/verify_codes.py \
    --combo-erp 组合装.csv --single-erp 单品.csv \
    --platform-file 快手店铺.xlsx --store 快手店铺 --output 核验结果.xlsx

Multiple --platform-file arguments are supported for batch testing. Normal workflow is
one store file per run. No source file is ever modified.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from typing import Any

import pandas as pd
from rapidfuzz import fuzz
import confirmed_mappings as confirmations
import sku_audit_reviews as sku_reviews
import confirmed_business_specs as business_specs
from quantity_syntax import packaging_products, without_rates, additive_quantities, unresolved_gifts
from quantity_syntax import quantity_prose, parenthetical_counts, apparel_component_text, TOTAL
from quantity_syntax import layered_packaging_details, numeric_piece_offer, without_remark_gifts
from manual_weight import apply_weight_guard
from confirmation_expiry import mark_expired
import agricultural_specs as agriculture
from erp_library import resolve_erp_inputs, DEFAULT_MANIFEST
from simple_report import write_simple_report

VERSION = "1.7"
PLATFORM_PROFILES: dict[str, Any] = {}
BUSINESS_RULES: dict[str, Any] = {}


def literal_text(value: Any) -> str:
    """Decode only ERP's explicit Excel text wrapper; never evaluate formulas."""
    raw = str(value if value is not None else "")
    match = re.fullmatch(r'="((?:[^"]|"")*)"', raw)
    return match.group(1).replace('""', '"') if match else raw

# ---------- Normalization / header mapping ----------
def norm_text(value: Any) -> str:
    if value is None:
        return ""
    s = unicodedata.normalize("NFKC", literal_text(value)).strip().lower()
    s = re.sub(r"[\s\u3000,，。;；:：!！?？、/\\|_+*×()（）\[\]【】{}<>《》\-—·“”‘’'\"`~]+", "", s)
    if s in {"nan", "none", "null"}:
        return ""
    return s


def norm_col(value: Any) -> str:
    s = norm_text(value)
    # pandas suffixes duplicate headers as .1, .2; preserve position separately.
    s = re.sub(r"\.\d+$", "", s)
    return s

HEADER_ALIASES = {
    "product_id": {"商品id", "商品编号", "itemid", "productid", "商品id编号", "id", "item_id", "product_id"},
    "title": {"商品标题", "商品名称", "商品名", "title", "itemtitle", "producttitle", "商品标题名称", "productname", "itemname", "商品短标题"},
    "sku_id": {"skuid", "sku编号", "规格id", "销售属性id", "variantid", "skuid编号", "sku_id", "规格编号", "商品规格id"},
    "sku_name": {"sku名称", "sku规格", "销售属性", "规格", "规格名称", "销售属性组合", "sku属性", "属性值", "sku名称规格", "规格属性", "商品规格", "商品规格名称", "规格描述", "规格值", "销售规格", "规格信息", "sku名称规格信息"},
    "code": {"skucode", "sku编码", "sku商家编码", "skumerchantcode", "skuskuouterid", "skuouterid", "外部sku编码", "外部编码", "商品编码", "商家编码", "outerid", "merchantcode", "编码", "货号", "商品货号", "款号"},
    "erp_item_id": {"erp货品编号", "货品编号", "erpitemid"},
    "spec_code": {"erp规格编码", "规格编码"},
    "quantity": {"规格数量"},
}


def load_alias_config(config_path: str | None = None) -> None:
    """Merge user-editable aliases from config/platform_aliases.json if present."""
    config_path = Path(config_path) if config_path else Path(__file__).resolve().parent.parent / "config" / "platform_aliases.json"
    if not config_path.exists():
        return
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
        PLATFORM_PROFILES.clear()
        PLATFORM_PROFILES.update(data.get("platforms", {}))
        BUSINESS_RULES.clear()
        BUSINESS_RULES.update(data.get("business_rules", {}))
        for canonical, aliases in data.get("canonical_fields", {}).items():
            if canonical in HEADER_ALIASES and isinstance(aliases, list):
                HEADER_ALIASES[canonical].update(norm_col(alias) for alias in aliases)
    except Exception as exc:
        raise RuntimeError(f"平台字段映射配置读取失败：{config_path}；{exc}")


load_alias_config()


def header_score(row: list[Any]) -> int:
    vals = [norm_col(x) for x in row]
    score = 0
    for canonical, aliases in HEADER_ALIASES.items():
        if any(v in aliases for v in vals):
            score += {"product_id": 2, "title": 3, "sku_id": 2, "sku_name": 3, "code": 3}.get(canonical, 0)
    return score


def read_csv_guess(path: str, **kwargs) -> pd.DataFrame:
    last_err = None
    for enc in ("utf-8-sig", "gb18030", "utf-8", "gbk"):
        try:
            return pd.read_csv(path, encoding=enc, dtype=str, keep_default_na=False, **kwargs)
        except Exception as exc:
            last_err = exc
    raise RuntimeError(f"无法读取 CSV：{path}；最后错误：{last_err}")


def detect_csv_header(path: str) -> tuple[pd.DataFrame, int]:
    preview = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with open(path, encoding=encoding, newline="") as stream:
                reader = csv.reader(stream)
                preview = [row for _, row in zip(range(12), reader)]
            break
        except UnicodeDecodeError:
            continue
    if preview is None:
        raise RuntimeError(f"无法读取CSV表头：{path}")
    best_row, best_score = 0, -1
    for i in range(len(preview)):
        score = header_score(preview[i])
        if score > best_score:
            best_row, best_score = i, score
    if best_score < 4:
        raise RuntimeError(f"无法自动识别平台 CSV 表头：{path}。请检查文件或补充字段映射。")
    return read_csv_guess(path, header=best_row), best_row


def detect_excel_sheet(path: str, explicit_sheet: str | None = None) -> tuple[pd.DataFrame, str, int]:
    xls = pd.ExcelFile(path)
    available_sheets = list(xls.sheet_names)
    xls.close()
    if explicit_sheet:
        if explicit_sheet not in available_sheets:
            raise RuntimeError(f"工作表不存在：{explicit_sheet}；可用工作表：{available_sheets}")
        sheets = [explicit_sheet]
    else:
        sheets = available_sheets
    candidates = []
    for sheet in sheets:
        preview = pd.read_excel(path, sheet_name=sheet, header=None, nrows=12, dtype=str, keep_default_na=False)
        best_row, best_score = 0, -1
        for i in range(len(preview)):
            score = header_score(preview.iloc[i].tolist())
            if score > best_score:
                best_row, best_score = i, score
        if best_score >= 4:
            # Prefer actual data-bearing sheets over hidden metadata sheets.
            try:
                body = pd.read_excel(path, sheet_name=sheet, header=best_row, dtype=str, nrows=250).fillna("")
                nonempty = int(body.astype(str).apply(lambda col: col.str.strip().ne("")).any(axis=1).sum())
            except Exception:
                nonempty = 0
            candidates.append((best_score * 1000 + min(nonempty, 999), best_score, nonempty, sheet, best_row))
    if not candidates:
        raise RuntimeError(f"无法在 Excel 中识别平台商品表头：{path}。请指定字段映射或检查文件。")
    candidates.sort(reverse=True)
    _, score, rows, sheet, header_row = candidates[0]
    # If equally plausible data sheets exist, fail safe instead of silently choosing one.
    ties = [c for c in candidates if c[2] > 0 and c[3] != sheet]
    if ties and not explicit_sheet:
        raise RuntimeError(f"发现多个可能的商品工作表：{[c[3] for c in candidates]}；请用 --sheet 指定。")
    df = pd.read_excel(path, sheet_name=sheet, header=header_row, dtype=str, keep_default_na=False)
    # A numeric cell formatted as 000000 is an identifier, not the number shown
    # by pandas. Recover explicit zero padding without changing source files.
    from openpyxl import load_workbook
    if Path(path).suffix.lower() != ".xls":
        book = load_workbook(path, read_only=True, data_only=True)
        worksheet = book[sheet]
        worksheet.reset_dimensions()
        identifier_fields = {"product_id", "sku_id", "code", "erp_item_id", "spec_code"}
        id_columns = [i for i, column in enumerate(df.columns) if any(norm_col(column) in HEADER_ALIASES[k] for k in identifier_fields)]
        for index, cells in enumerate(worksheet.iter_rows(min_row=header_row + 2, max_row=header_row + 1 + len(df), max_col=len(df.columns))):
            for column in id_columns:
                cell = cells[column]
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and float(cell.value).is_integer() and re.fullmatch(r"0+", cell.number_format or ""):
                    df.iat[index, column] = str(int(cell.value)).zfill(len(cell.number_format))
        book.close()
    return df, sheet, header_row


def choose_platform_columns(df: pd.DataFrame, platform: str | None = None) -> dict[str, str | None]:
    cols = list(df.columns)
    normalized = [(c, norm_col(c), i) for i, c in enumerate(cols)]

    def first_match(canonical: str) -> str | None:
        explicit = PLATFORM_PROFILES.get(platform, {}).get("columns", {}).get(canonical)
        if explicit is not None:
            if explicit not in cols:
                raise RuntimeError(f"平台{platform}配置字段不存在：{explicit}；请确认字段映射。")
            return explicit
        aliases = HEADER_ALIASES[canonical]
        hits = [(c, n, i) for c, n, i in normalized if n in aliases]
        if not hits:
            return None
        if canonical == "code":
            # SKU-level code fields outrank item-level merchant-code fields.
            def rank(item):
                c, n, i = item
                if n in {"skucode", "sku编码", "sku商家编码", "skumerchantcode", "skuskuouterid", "skuouterid", "外部sku编码"}:
                    return (100, i)
                if "商家编码" in n or n == "merchantcode":
                    return (60, i)  # Equal semantic ranks require an explicit mapping.
                if n == "outerid":
                    return (40, i)
                return (20, i)
            best_rank = max(rank(item)[0] for item in hits)
            hits = [item for item in hits if rank(item)[0] == best_rank]
        if len(hits) != 1:
            raise RuntimeError(f"字段{canonical}有多个可能列：{[h[0] for h in hits]}；请在平台配置中明确映射。")
        return hits[0][0]

    result = {key: first_match(key) for key in HEADER_ALIASES}
    # Tmall-like templates can have multiple sales-attribute columns; concatenate them if necessary.
    if result["sku_name"] is None:
        attr_cols = [c for c, n, _ in normalized if n in {"属性对", "销售属性", "销售属性值", "sku属性值"}]
        if attr_cols:
            result["sku_name"] = attr_cols[-1]
    required = ["product_id", "title", "code"]
    missing = [k for k in required if result.get(k) is None]
    if missing:
        raise RuntimeError(f"无法识别必要字段 {missing}；实际字段为：{cols}。需要人工确认字段映射。")
    if result.get("erp_item_id") == result.get("code") or result.get("spec_code") == result.get("code"):
        raise RuntimeError("商品身份标识不能映射到当前编码列。")
    return result

# ---------- ERP database ----------
def load_erp(combo_path: str, single_path: str) -> pd.DataFrame:
    parts = []
    for path, origin, name_col in ((combo_path, "组合装", "组合装名称"), (single_path, "单品", "货品名称")):
        raw = read_csv_guess(path)
        if "商家编码" not in raw or name_col not in raw:
            raise RuntimeError(f"{origin}ERP缺少商家编码或{name_col}；实际字段：{list(raw.columns)}")
        rows = raw.copy()
        rows["原始字段JSON"] = [json.dumps(r, ensure_ascii=False) for r in raw.to_dict("records")]
        rows["ERP来源"] = origin
        rows["ERP源文件"] = str(Path(path).resolve())
        rows["ERP文件SHA256"] = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        rows["原始行号"] = range(2, len(rows) + 2)
        rows["ERP商家编码"] = raw["商家编码"].map(literal_text).str.strip()
        rows["ERP商品名称"] = raw[name_col].map(literal_text).str.strip()
        for canonical, original in (("ERP类别", "类别"), ("ERP分类", "分类"), ("ERP重量", "重量"),
                                    ("ERP货品编号", "货品编号"), ("ERP规格编码", "规格编码"), ("ERP规格数量", "规格数量")):
            rows[canonical] = raw[original].map(literal_text).str.strip() if original in raw else ""
        parts.append(rows)
    merged = pd.concat(parts, ignore_index=True).fillna("")
    merged["_name_norm"] = merged["ERP商品名称"].map(norm_text)
    return merged


COUNT_UNITS = r"片|对|包|袋|只|双|个|件|盒|副|套"
STRONG_PACKAGING_TERMS = ('外袋','内袋','卷材','纸板','纸箱','外箱','包装箱','包装材料','暖贴包装')
SOFT_PACKAGING_TERMS = ('彩盒','彩箱','纸盒','内盒','包装袋')


def _scalar_quantity_details(value: str) -> tuple[tuple[int, str] | None, list[str]]:
    s = unicodedata.normalize("NFKC", literal_text(value)).lower()
    issues = []
    totals = list(dict.fromkeys((int(n), u) for n, u in re.findall(rf"(?:{TOTAL})\s*(\d+)\s*({COUNT_UNITS})", s)))
    # Adjacent outer/inner totals state the physical count, not a conversion.
    nested = re.search(rf'(?:{TOTAL})\s*\d+\s*(?:盒|包|袋)\s*(?:共|合计|总计)?\s*\(?\s*(\d+)\s*(片|对|只|个)',s)
    if nested:
        totals = [q for q in totals if q[1] not in {'盒','包','袋'}] + [(int(nested[1]),nested[2])]
        totals = list(dict.fromkeys(totals))
    pack_totals, pack_issues = packaging_products(s, COUNT_UNITS)
    rates = re.findall(r'(\d+)\s*(片|对|只|个)\s*[/／]\s*(盒|包|袋)',s)
    if len(rates) == 1:
        per,unit,outer = rates[0]
        outer_counts = list(dict.fromkeys(int(n) for n in re.findall(rf'(\d+)\s*{outer}',s)))
        if len(outer_counts) == 1:
            pack_totals.append((int(per)*outer_counts[0],unit))
    issues.extend(pack_issues)
    ingredients = {x for x in ("艾草", "益母草", "生姜") if x in s}
    each = re.search(rf"各\s*(\d+)\s*({COUNT_UNITS})", s)
    each_total = (int(each.group(1)) * 3, each.group(2)) if each and len(ingredients) == 3 else None
    if each and not each_total:
        issues.append("各数量的商品组成不明确")
    formula_counts = re.findall(r"(?:艾草|益母草|生姜)\s*(\d+)(?:片)?", s)
    formula_total = (sum(map(int, formula_counts)), "片") if len(formula_counts) == 3 else None
    direct = list(dict.fromkeys((int(n), u) for n, u in re.findall(rf"(\d+)\s*({COUNT_UNITS})", without_rates(s))))
    gift = re.search(rf"(?:买|单拍)?\s*(\d+)\s*({COUNT_UNITS})?\s*(?:送|赠)\s*(\d+)\s*({COUNT_UNITS})?", s)
    gift_total = None
    if gift:
        first_unit, second_unit = gift[2], gift[4]
        unit = second_unit or first_unit
        different_product = re.match(r'足贴|鞋垫|暖贴|眼罩|护膝|护腰|暖手|精油|收纳',s[gift.end():].lstrip())
        if unit and (not first_unit or not second_unit or first_unit == second_unit) and not different_product:
            gift_total = (int(gift[1]) + int(gift[3]), unit)
    if any(len({n for n,unit in totals if unit==u})>1 for _,u in totals):
        issues.append("明示总数量相互冲突")
    preferred = next((q for q in totals if q[1] in {'片','只','个'}), totals[0] if totals else None)
    quantity = preferred or each_total or (pack_totals[0] if len(pack_totals) == 1 else None) or gift_total
    if quantity is None and len(direct) == 1 and not each:
        quantity = direct[0]
    if quantity is None and formula_total:
        quantity = formula_total
    # 1包30只 denotes one outer pack; never use an arbitrary first count.
    if quantity is None and re.search(rf"1\s*(?:包|盒|袋)\s*\d+\s*(片|对|只|个)", s):
        units = [(n, u) for n, u in direct if u not in {"包", "盒", "袋"}]
        quantity = units[0] if len(units) == 1 else None
    for other in pack_totals + ([each_total] if each_total else []) + ([formula_total] if formula_total else []) + ([gift_total] if gift_total else []):
        if quantity and other[1] == quantity[1] and other[0] != quantity[0]:
            issues.append("总数与包装倍数或配方组成不一致")
    if re.search(r"送|赠", s) and "同款" not in s:
        issues.append("赠品或主赠品数量未建立对应关系")
    if ("+" in s or "＋" in s) and not formula_total and not (len(ingredients) == 3 and len(direct) == 1):
        issues.append("套装组成未建立结构化对应关系")
    if quantity is None:
        issues.append("数量未明确或多个数量表达无法确定总数")
    if pack_issues:
        quantity = None
    return quantity, list(dict.fromkeys(issues))


def quantity_structure(value: str) -> dict[str, Any]:
    detail = deepcopy(_cached_quantity_structure(literal_text(value), parsing_rule_signature()))
    text = unicodedata.normalize('NFKC', literal_text(value))
    layers=layered_packaging_details(text.lower())
    if layers:
        detail['packaging_details']=layers
    if '三合一' in text or all(name in text for name in ('艾草','生姜','益母草')):
        detail['recipe_details'] = [{'formula':name,'quantity':(int(n),'片')}
            for name,n in re.findall(r'(艾草|生姜|益母草)\s*(\d+)(?:片)?',text)]
        detail['comparison_scope'] = 'total_only'
    return detail


def parsing_rule_signature():
    return (tuple(sorted((k, tuple(v)) for k,v in BUSINESS_RULES.get('family_terms', {}).items())),
            tuple(BUSINESS_RULES.get('pair_piece_families', [])))


@lru_cache(maxsize=32768)
def _cached_quantity_structure(value, rule_signature):
    text = unicodedata.normalize('NFKC', literal_text(value)).lower()
    text = quantity_prose(text)
    main_text = without_remark_gifts(text)
    oil_remark = main_text!=text
    text = main_text
    text, breakdown_issues, breakdown_details = parenthetical_counts(text, COUNT_UNITS)
    text, offer = numeric_piece_offer(text,extract_variant_groups)
    parsed = additive_quantities(text, _scalar_quantity_details, extract_families,
        extract_variant_groups, COUNT_UNITS, BUSINESS_RULES.get('pair_piece_families', []))
    if parsed is not None:
        parsed['raw'] = literal_text(value)
        if offer:
            parsed['offer_details']=offer
            if offer['issues']:
                parsed.update(qty=None,incomplete=True,issues=[*parsed['issues'],*offer['issues']])
        if oil_remark:parsed['remark_gifts']=['精油']
        if breakdown_issues:
            parsed.update(qty=None, incomplete=True, issues=[*parsed['issues'], *breakdown_issues])
        parsed['breakdown_details'] = breakdown_details
        return parsed
    qty, issues = _scalar_quantity_details(text)
    equivalent = re.search(r'(\d+)\s*片(?:装)?\s*[/／]\s*(\d+)\s*对',text)
    equivalent_issue = False
    alternatives = list(dict.fromkeys((int(n),u) for n,u in re.findall(rf'(?:{TOTAL})\s*(\d+)\s*({COUNT_UNITS})',text)))
    if equivalent:
        permitted = bool(extract_families(text) & set(BUSINESS_RULES.get('pair_piece_families', [])))
        qty=(int(equivalent[1]),'片')
        alternatives += [(int(equivalent[1]),'片'),(int(equivalent[2]),'对')]
        issues=[i for i in issues if i!='数量未明确或多个数量表达无法确定总数']
        if permitted and int(equivalent[1]) != 2*int(equivalent[2]):
            equivalent_issue=True
            issues.append('已授权品类的片/对明示数量互相矛盾')
    gift_issues = unresolved_gifts(text, COUNT_UNITS, extract_families, _scalar_quantity_details)
    issues = list(dict.fromkeys([*issues, *gift_issues, *breakdown_issues]))
    inconsistent = any(issue in {'明示总数量相互冲突','总数与包装倍数或配方组成不一致'} for issue in issues)
    incomplete = bool(gift_issues or breakdown_issues or equivalent_issue or inconsistent) or any('包装倍数' in issue for issue in issues)
    if qty is None and re.search(r'\d+\s*(?:片|对|只|个)\s*[/／]\s*(?:盒|包|袋)',text):
        incomplete = True
        issues = [*issues, '每包数量已知但包装数或总数未明确']
    if '三合一' in text and qty is None:
        incomplete = True
        issues = [*issues, '三合一总数量未明确']
    if offer and offer['issues']:
        incomplete=True
        issues=[*issues,*offer['issues']]
    if oil_remark and qty is None:
        incomplete=True
        issues=[*issues,'精油仅作备注，主商品数量尚未明确']
    result={'raw': literal_text(value), 'qty': None if incomplete else qty, 'parsed_total':qty,
            'parts': [], 'issues': issues,
            'incomplete': incomplete, 'additive': False, 'breakdown_details':breakdown_details,
            'alternative_quantities':alternatives}
    if offer: result['offer_details']=offer
    if oil_remark: result['remark_gifts']=['精油']
    return result


def quantity_details(value: str) -> tuple[tuple[int, str] | None, list[str]]:
    parsed = quantity_structure(value)
    return parsed['qty'], parsed['issues']


def extract_quantity(text: str, families: set[str] | None = None) -> tuple[int, str] | None:
    return quantity_details(text)[0]


def extract_families(text: str) -> set[str]:
    segments = [norm_text(piece) for piece in re.split(r"[\s\u3000,，。;；:：!！?？、/\\|_+*×()（）\[\]【】{}<>《》\-—·“”‘’'\"`~]+", literal_text(text))]
    segments = [piece for piece in segments if piece]
    if not segments:
        return set()
    contains = lambda word: any(norm_text(word) in piece for piece in segments)
    # Packaging material takes precedence over the product that the packaging is for.
    if any(contains(x) for x in STRONG_PACKAGING_TERMS):
        return {"包装材料"}
    families = set()
    for family, terms in BUSINESS_RULES.get("family_terms", {}).items():
        if any(contains(term) for term in terms):
            families.add(family)
    if any(contains(x) for x in ("除湿袋", "干燥剂", "吸湿", "防潮盒")):
        families.add("除湿用品")
    if contains("手套"):
        families.add("手套")
    if contains("眼罩"):
        families.add("眼罩")
    if contains("鞋垫"):
        families.add("发热鞋垫")
    if any(contains(x) for x in ("暖手宝", "暖手贴")):
        families.add("暖手用品")
    if any(contains(x) for x in ("暖足贴", "足贴")):
        families.add("足贴")
    if any(contains(x) for x in ("暖贴", "暖宝宝", "暖身贴", "宫暖贴", "热敷贴", "温暖贴")):
        families.add("暖贴")
    if not families and any(contains(x) for x in SOFT_PACKAGING_TERMS):
        return {'包装材料'}
    return families


def extract_variant_groups(text: str) -> dict[str, set[str]]:
    s = norm_text(text)
    groups: dict[str, set[str]] = defaultdict(set)
    if not s:
        return groups
    # Formula / ingredient group.
    ingredients = {x for x in ("艾草", "益母草", "生姜") if x in s}
    if "三合一" in s or {"艾草", "益母草", "生姜"}.issubset(ingredients):
        groups["formula"].add("三合一")
    else:
        groups["formula"].update(ingredients)
    # Mutually exclusive product types where explicitly stated.
    if "增强" in s:
        groups["heat_type"].add("增强")
    if "通用" in s or "通用性" in s:
        groups["heat_type"].add("通用")
    # Product style / brand family aliases.
    for token in ("卡通", "hellokitty", "kt", "凯蒂猫"):
        if token in s:
            groups["style"].add(token)
    for key, terms in {
        "国潮": ("国潮",), "哪吒": ("哪吒",), "小魔仙": ("小魔仙",), "暖暖": ("暖暖",),
        "马蹄型": ("马蹄型", "马蹄形"), "可穿戴": ("可穿戴",),
    }.items():
        if any(t in s for t in terms):
            groups["style"].add(key)
    if "男款" in s or "男码" in s or "男士" in s:
        groups["gender"].add("男款")
    if "女款" in s or "女码" in s or "女士" in s:
        groups["gender"].add("女款")
    if "儿童" in s or "童款" in s or "儿童款" in s:
        groups["gender"].add("儿童款")
    for material in ("pvc", "cpe", "丁腈", "硅胶", "无纺布"):
        if material in s:
            groups["material"].add(material)
    groups["size"].update(re.findall(r"(?<![a-z])(?:[2-5]?x{1,4}l|[sml])(?![a-z])", s))
    for color in ("粉", "蓝", "黑", "白", "灰", "黄", "紫", "茶", "红", "透明"):
        if color in s:
            groups["color"].add(color)
    return groups


def build_features(title: str, sku_name: str, erp_name: str = "", erp_category: str = "") -> dict[str, Any]:
    is_erp = bool(erp_name)
    specific = literal_text(erp_name if is_erp else sku_name)
    title = literal_text(title)
    combined = specific if is_erp else f"{title} {specific}"
    families = extract_families(specific)
    family_source = 'sku' if families else 'unknown'
    if not families:
        fallback = extract_families(erp_category if is_erp else title)
        # A selected SKU has priority. SEO titles listing multiple categories
        # are candidate evidence, not proof of a category contradiction.
        shoe_like_spec = re.search(r'\d{2}\s*[-~至]\s*\d{2}\s*码|男码|女码|儿童码',specific)
        if is_erp or (len(fallback) == 1 and not (fallback == {'足贴'} and shoe_like_spec)):
            families = fallback
            family_source = 'erp_category' if is_erp else 'title'
    component_text = apparel_component_text(title, unicodedata.normalize('NFKC',specific).lower()) if not is_erp else None
    detail = quantity_structure(component_text or specific)
    if component_text:
        detail['raw'] = specific
        detail['contextual_main'] = '标题唯一护膝/护腰品类与SKU明示尺码；未推定衣物件数'
    quantity, issues = detail['qty'], detail['issues']
    # SKU attributes take precedence. Title-only attributes generate candidates;
    # they do not fill missing SKU specifications for automatic decisions.
    variants = extract_variant_groups(specific)
    result = {"text": combined, "quantity_text": specific, "norm": norm_text(combined), "families": families,
            "qty": quantity, "qty_issues": issues, "variants": variants,
            "brand_warmhuo": "暖火" in combined, "is_erp": is_erp,
            "erp_item_id": "", "spec_code": ""}
    result['title_text'] = title
    result['family_source'] = family_source
    result['quantity_structure'] = detail
    if component_text:
        result['component_parse_text'] = component_text
    result['packaging_issue'] = ('彩盒等包装名称需确认是成品还是包装材料'
        if any(w in specific for w in SOFT_PACKAGING_TERMS) and not any(w in specific for w in STRONG_PACKAGING_TERMS) else '')
    result['agricultural'] = agriculture.parse_spec(specific, title if not is_erp else '', erp_category)
    result["evidence"] = structured_evidence(specific, title if not is_erp else "", erp_category, result)
    return result


def structured_evidence(specific: str, title: str, category: str, features: dict[str, Any]) -> dict[str, set[str]]:
    """Extract literal attributes, without guessing brand/size aliases."""
    s = unicodedata.normalize("NFKC", specific).lower()
    product = unicodedata.normalize("NFKC", title + " " + category).lower()
    attrs = {k: set(v) for k, v in features["variants"].items() if v}
    # Usage prose such as 适合男士/女士 is not a selected SKU gender.
    attrs["gender"] = {label for label, words in (("男款", ("男款", "男码")), ("女款", ("女款", "女码")), ("儿童款", ("儿童款", "儿童码", "童款"))) if any(w in s for w in words)}
    for group, terms in {
        "brand": ("暖火", "kt", "hellokitty", "凯蒂猫", "小黄鸭", "哪吒", "小魔仙"),
        "scent": ("薰衣草", "玫瑰", "桂花", "洋甘菊", "无香", "四合一"),
        "eye_type": ("蒸汽", "冰凉", "冷敷"),
        "shape": ("u型", "圆形", "长方形", "马蹄型"),
    }.items():
        values = {term for term in terms if term in s}
        if not values and group in {"brand", "eye_type"}:
            inherited = {term for term in terms if term in product}
            values = inherited if len(inherited) == 1 else set()
        attrs[group] = values
    if not attrs.get("material"):
        materials = {m for m in ("pvc", "cpe", "丁腈", "硅胶") if m in product}
        attrs["material"] = materials if len(materials) == 1 else set()
    attrs["packaging"] = {unit for unit in ("盒", "袋", "桶", "箱") if re.search(rf"(?:{unit}装|\d+\s*{unit})", s)}
    attrs["dimensions"] = {f"{a}-{b}码" for a, b in re.findall(r"(\d{2})\s*[-/~]\s*(\d{2})\s*码", s)}
    attrs["capacity"] = {format((Decimal(n) * (1000 if unit in {"l", "升"} else 1)).normalize(), "f") + "ml" for n, unit in re.findall(r"(\d+(?:\.\d+)?)\s*(ml|毫升|l|升)(?![a-z])", s)}
    attrs["model"] = set(re.findall(r"型号\s*[:：]?\s*([a-z0-9-]+)", s))
    attrs["duration"] = {n + "小时" for n in re.findall(r"(?:发热|持续|持久)?\s*(\d+)\s*(?:小时|h)(?![a-z])", s)}
    attrs["series"] = {token for token in ("经典", "升级", "增强") if token in s}
    attrs["formula_composition"] = {name + ":" + amount + "片" for name, amount in re.findall(r"(艾草|益母草|生姜)\s*(\d+)(?:片)?", s)}
    each = re.search(r"各\s*(\d+)\s*片", s)
    if each and all(name in s for name in ("艾草", "益母草", "生姜")):
        attrs["formula_composition"] = {name + ":" + each.group(1) + "片" for name in ("艾草", "益母草", "生姜")}
    return {k: v for k, v in attrs.items() if v}


def structured_match(pf: dict[str, Any], records: list[dict[str, Any]]) -> tuple[dict[str, Any] | None, int, str]:
    labels = {"brand": "品牌", "formula": "配方", "heat_type": "发热类型", "style": "款式",
              "material": "材质", "size": "尺码", "scent": "香型", "eye_type": "眼罩类型",
              "capacity": "容量", "shape": "形状/型号", "color": "颜色", "model": "型号",
              "series": "系列", "dimensions": "数字尺码范围", "packaging": "包装", "gender": "男女/儿童款", "duration": "发热时长", "formula_composition": "三种配方各自数量"}
    families = pf["families"]
    if len(families) != 1:
        return None, 0, "商品品类不明确或包含多个品类"
    family = next(iter(families))
    profile = BUSINESS_RULES.get("structured_profiles", {}).get(family)
    if profile is None:
        return None, 0, "该品类尚无完整结构化身份规则"
    attrs = pf.get("evidence", {})
    required = profile["required"]
    missing = [field for field in required if not attrs.get(field)]
    if missing:
        return None, 0, "平台身份属性缺少：" + "/".join(labels.get(k, k) for k in missing)
    if attrs.get("formula") == {"三合一"} and len(attrs.get("formula_composition", set())) != 3:
        return None, 0, "三合一合计数量已知，但三种配方各自数量尚未确认"
    min_fields = profile.get("min_identity_fields", len(required))
    identity_fields = {"brand", "formula", "heat_type", "style", "material", "size", "scent", "eye_type", "capacity", "shape", "model", "series"}
    if sum(bool(attrs.get(k)) for k in identity_fields) < min_fields:
        return None, 0, "平台缺少足够的独立身份属性，名称或数量不足以判定"
    possible = []
    for record in records:
        ef = record["features"]
        other = ef.get("evidence", {})
        if ef["families"] != families or any(other.get(k) and other[k] != attrs[k] for k in required):
            continue
        # Missing attributes/quantities remain possible alternatives. Do not
        # obtain uniqueness by dropping an incomplete ERP record.
        if any(attrs[k] != other[k] for k in attrs.keys() & other.keys()):
            continue
        if quantity_compare(pf["qty"], ef["qty"], families) is False:
            continue
        possible.append(record)
    if not possible:
        return None, 0, "未找到满足必要身份属性及现有规格的ERP记录"
    if len(possible) != 1:
        return None, len(possible), f"结构化身份下存在{len(possible)}条可能ERP记录，未唯一"
    chosen = possible[0]
    ef = chosen["features"]
    other = ef.get("evidence", {})
    if attrs != other:
        return None, 1, "平台与ERP明示属性不完整：" + "/".join(labels.get(k, k) for k in sorted(attrs.keys() ^ other.keys()))
    quantity_optional = profile.get("count_optional", False) and pf["qty"] is None and ef["qty"] is None
    if not quantity_optional and quantity_compare(pf["qty"], ef["qty"], families) is not True:
        return None, 1, "规格数量不完整或未一致"
    issues = pf.get("qty_issues", []) + ef.get("qty_issues", [])
    if quantity_optional:
        issues = [issue for issue in issues if issue != "数量未明确或多个数量表达无法确定总数"]
    if issues:
        return None, 1, "；".join(dict.fromkeys(issues))
    return chosen, 1, "品类必要身份属性、全部明示规格及数量一致；全库唯一，未使用名称分数"


def three_in_one_total_match(pf, current, records, approved):
    """Apply the user's quantity-only policy to identified three-in-one products."""
    def is_mixed(features):
        return (features.get("evidence", {}).get("formula") == {"三合一"}
                and len(features["families"]) == 1 and "包装材料" not in features["families"])

    if not is_mixed(pf):
        return None, 0, "三合一混合装品类不明确或含其他商品，仍需确认"
    if approved:
        possible = [approved["record"]]
    elif current:
        possible = current
    else:
        possible = [r for r in records if is_mixed(r["features"])
                    and r["features"]["families"] == pf["families"]
                    and quantity_compare(pf["qty"], r["features"]["qty"], pf["families"]) is not False]
    if len(possible) != 1:
        return None, len(possible), f"三合一总数量规则下存在{len(possible)}条可能ERP记录，无法唯一确定编码"
    chosen = possible[0]
    ef = chosen["features"]
    if not is_mixed(ef):
        return None, 1, "当前或已确认ERP记录为非三合一混合装，数量相同不能据此通过"
    if ef["families"] != pf["families"]:
        return None, 1, "平台与ERP三合一商品品类不一致"
    if quantity_compare(pf["qty"], ef["qty"], pf["families"]) is not True:
        return None, 1, "三合一总数量缺失、不一致或单位无法换算"
    issues = []
    for features in (pf, ef):
        for issue in features.get("qty_issues", []):
            # A stated received total covers the numeric buy/gift offer. It
            # does not establish an unrelated named gift's product relation.
            quantity_text = without_remark_gifts(features.get("quantity_text", features["text"]))
            numeric_offer = re.search(r"买\s*(\d+)\s*片?\s*送\s*(\d+)\s*片?\s*[【\[]?\s*(?:到手|共|合计|总计)\s*(\d+)\s*片", quantity_text)
            if (issue == "赠品或主赠品数量未建立对应关系" and numeric_offer
                    and not re.search(r"送|赠", quantity_text[:numeric_offer.start()] + quantity_text[numeric_offer.end():])):
                bought, gifted, received = map(int, numeric_offer.groups())
                if bought + gifted != received:
                    issues.append("买赠数量之和与明示到手总数不一致")
                continue
            issues.append(issue)
    if issues:
        return None, 1, "三合一总数量仍有问题：" + "；".join(dict.fromkeys(issues))
    return chosen, 1, "按用户确认规则：双方为同品类三合一混合装，总数量一致；配比不参与核验"


def quantity_compare(a: tuple[int, str] | None, b: tuple[int, str] | None, families: set[str]) -> bool | None:
    if not a or not b:
        return None
    av, au = a; bv, bu = b
    if au == bu:
        return av == bv
    if families & set(BUSINESS_RULES.get("pair_piece_families", [])) and {au, bu}.issubset({"片", "对"}):
        a_pieces = av if au == "片" else av * 2
        b_pieces = bv if bu == "片" else bv * 2
        return a_pieces == b_pieces
    return None


def feature_conflicts(pf: dict[str, Any], ef: dict[str, Any]) -> list[str]:
    conflicts = []
    p_fams, e_fams = pf["families"], ef["families"]
    if p_fams and e_fams and not (p_fams & e_fams):
        conflicts.append("商品类型不一致")
    qcmp = quantity_compare(pf["qty"], ef["qty"], p_fams | e_fams)
    if qcmp is False:
        conflicts.append("规格数量不一致")
    pvars, evars = pf["variants"], ef["variants"]
    for group in ("formula", "heat_type", "gender", "material", "size", "color"):
        pset, eset = pvars.get(group, set()), evars.get(group, set())
        if pset and eset and not (pset & eset):
            conflicts.append({"formula": "成分/配方不一致", "heat_type": "产品类型不一致", "gender": "性别/尺码类型不一致", "material": "材质不一致", "size": "尺码不一致", "color": "颜色不一致"}[group])
    # Style family can be multiple (e.g. KT + classic); only conflict on exclusive named styles.
    exclusive_styles = {"卡通", "hellokitty", "kt", "凯蒂猫", "国潮", "哪吒", "小魔仙", "暖暖"}
    pstyle = pvars.get("style", set()) & exclusive_styles
    estyle = evars.get("style", set()) & exclusive_styles
    if pstyle and estyle and not (pstyle & estyle):
        conflicts.append("产品系列/款式不一致")
    return conflicts


def identity_score(pf: dict[str, Any], ef: dict[str, Any]) -> tuple[float, list[str], bool]:
    # Scores rank suggestions only. Even identical names cannot authorize a code.
    fuzzy = fuzz.WRatio(pf["norm"], ef["norm"]) if pf["norm"] and ef["norm"] else 0
    score = fuzzy * 0.6
    evidence = [f"候选名称相似度{fuzzy:.0f}（不作为自动判定）"]
    if pf["families"] & ef["families"]:
        score += 20
        evidence.append("品类词面相关")
    if quantity_compare(pf["qty"], ef["qty"], pf["families"] | ef["families"]) is True:
        score += 10
        evidence.append("明示数量一致（含已确认单位换算）")
    if any(pf["variants"].get(g, set()) & ef["variants"].get(g, set()) for g in pf["variants"]):
        score += 10
    evidence.extend(feature_conflicts(pf, ef))
    return round(min(100, score), 2), evidence, False


def candidate_rows(pf: dict[str, Any], erp_records: list[dict[str, Any]], top_n: int | None = 3) -> list[dict[str, Any]]:
    scored = []
    for er in erp_records:
        ef = er["features"]
        related_family = bool(pf["families"] & ef["families"])
        # A different recognized product family is not a plausible substitute.
        if pf["families"] and ef["families"] and not related_family:
            continue
        fuzzy = fuzz.WRatio(pf["norm"], ef["norm"]) if pf["norm"] and ef["norm"] else 0
        if not related_family and fuzzy < 60:
            continue
        if not ef["norm"]:
            continue
        score, evidence, _ = identity_score(pf, ef)
        scored.append({"record": er, "score": score, "evidence": evidence,
                       "strong": False, "conflicts": feature_conflicts(pf, ef)})
    scored.sort(key=lambda x: (-x["score"], x["record"]["code"], x["record"]["name"], x["record"]["source"]))
    return scored if top_n is None else scored[:top_n]


def classify_row(current_code: str, pf: dict[str, Any], erp_by_code: dict[str, list[dict[str, Any]]], erp_records: list[dict[str, Any]]) -> dict[str, Any]:
    code = literal_text(current_code).strip()
    current = erp_by_code.get(code, []) if code else []
    candidates = candidate_rows(pf, erp_records, top_n=None)
    identity = literal_text(pf.get("erp_item_id", "")).strip()
    spec_code = literal_text(pf.get("spec_code", "")).strip()
    exact = [r for r in erp_records if identity and r.get("item_id", "") == identity
             and (not spec_code or r.get("spec_code", "") == spec_code)]
    result = {"核验状态": "待人工确认", "应改为编码": "", "匹配依据": "", "核验说明": "",
              "ERP数据来源": "", "ERP唯一对应编码": "", "候选总数": len(candidates),
              "独立身份命中条数": len(exact), "同规格候选数": sum(not c["conflicts"] and quantity_compare(pf["qty"],c["record"]["features"]["qty"],pf["families"]|c["record"]["features"]["families"]) is True for c in candidates),
              "当前编码ERP名称": " | ".join(r["name"] for r in current),
              "当前编码ERP来源": " | ".join(r["source"] for r in current),
              "数量解析": str(pf["qty"] or "未明确"), "规格问题": "；".join(pf.get("qty_issues", []))}
    related_sources = sorted({r["source"] for r in current} | {c["record"]["source"] for c in candidates})
    if related_sources:
        result["ERP数据来源"] = "待核候选：" + " | ".join(related_sources)
    for i in range(3):
        suffix = "" if i == 0 else str(i+1)
        candidate = candidates[i] if i < len(candidates) else None
        for field, key in (("名称","name"), ("编码","code"), ("来源","source")):
            result[f"ERP候选{suffix}{field}"] = candidate["record"][key] if candidate else ""
        result[f"ERP候选{suffix}分数"] = candidate["score"] if candidate else ""
        result[f"ERP候选{suffix}冲突"] = "；".join(candidate["conflicts"] + candidate["record"]["features"].get("qty_issues", [])) if candidate else ""
    blockers = []
    if pf.get('packaging_issue') or any(r['features'].get('packaging_issue') for r in current):
        blockers.append('成品或包装材料身份尚未确认')
    if len(current) > 1:
        blockers.append("ERP合并后当前编码对应多条记录")
    if pf.get("duplicate_sku"):
        blockers.append("平台SKU ID重复，未删除记录")
    chosen, mode = None, ""
    structured_count, structured_reason = 0, ""
    approved = pf.get("confirmation")
    if pf.get("confirmation_issue"):
        blockers.append(pf["confirmation_issue"])
    elif BUSINESS_RULES.get("three_in_one_mixed_total_only") and pf.get("evidence", {}).get("formula") == {"三合一"}:
        chosen, structured_count, structured_reason = three_in_one_total_match(pf, current, erp_records, approved)
        mode = "三合一混合装总数量（用户规则）" if chosen else ""
        if not chosen:
            blockers.append(structured_reason)
    elif approved:
        chosen, mode = approved["record"], "人工确认映射复用"
        if identity and (chosen.get("item_id") != identity or (spec_code and chosen.get("spec_code") != spec_code)):
            blockers.append("人工映射与独立ERP身份字段冲突")
    elif identity:
        mode = "独立ERP身份编号"
        if len(exact) == 1:
            chosen = exact[0]
        else:
            blockers.append(f"独立货品编号/规格编码命中{len(exact)}条ERP记录，未唯一")
    else:
        chosen, structured_count, structured_reason = structured_match(pf, erp_records)
        mode = "完整结构化身份与规格" if chosen else ""
        if not chosen:
            blockers.append(structured_reason)
    if chosen:
        ef = chosen["features"]
        if mode == "独立ERP身份编号":
            blockers.extend(feature_conflicts(pf, ef))
            blockers.extend(pf.get("qty_issues", []))
            blockers.extend(ef.get("qty_issues", []))
            if quantity_compare(pf["qty"], ef["qty"], pf["families"] | ef["families"]) is not True:
                blockers.append("数量缺失、冲突或单位未授权换算")
            for group in set(pf["variants"]) | set(ef["variants"]):
                if pf["variants"].get(group, set()) != ef["variants"].get(group, set()):
                    blockers.append(f"{group}规格属性不完整或不一致")
        if not chosen["code"] or len(erp_by_code.get(chosen["code"], [])) != 1:
            blockers.append("拟对应ERP编码为空或在合并库中重复")
    result.update({"判定路径": mode if chosen and not blockers else "",
                   "结构化可能ERP数": structured_count, "结构化核验说明": structured_reason,
                   "平台结构化属性": json.dumps({k:sorted(v) for k,v in pf.get("evidence", {}).items()}, ensure_ascii=False, sort_keys=True),
                   "对应ERP结构化属性": json.dumps({k:sorted(v) for k,v in chosen["features"].get("evidence", {}).items()}, ensure_ascii=False, sort_keys=True) if chosen else "",
                   "人工确认人": approved["confirmed_by"] if approved else "",
                   "人工确认依据": approved["basis"] if approved else ""})
    if chosen and not blockers:
        result["ERP数据来源"] = chosen["source"]
        result["ERP唯一对应编码"] = chosen["code"]
        if mode == "三合一混合装总数量（用户规则）":
            result["规格问题"] = ""
        result["匹配依据"] = mode + "；" + (structured_reason if mode in {"完整结构化身份与规格", "三合一混合装总数量（用户规则）"} else "商品关系及对应规格已验证，ERP编码全库唯一")
        if not code:
            result["核验状态"] = "编码缺失（待补录）"
            result["核验说明"] = "当前编码为空；应改为编码留空，对应编码列供人工补录核对。"
        elif code == chosen["code"]:
            result["核验状态"] = "编码正确"
            result["核验说明"] = "三合一混合装总数量一致，按用户规则核对通过；不核实配比。" if mode == "三合一混合装总数量（用户规则）" else "商品身份与对应规格已验证，编码一致。"
        else:
            result["核验状态"] = "编码错误（建议修改）"
            result["应改为编码"] = chosen["code"]
            result["核验说明"] = "商品身份与对应规格已验证，当前非空编码不同。"
    else:
        if not candidates and not current and not exact and not pf.get("duplicate_sku"):
            result["核验状态"] = "ERP无匹配"
        result["匹配依据"] = "；".join(dict.fromkeys(blockers))
        notes = ["当前编码为空"] if not code else []
        if candidates:
            notes.append(f"全库检索{len(candidates)}条词面候选，仅展示前3条；候选分数不授权更码")
        else:
            notes.append("未检索到相关候选，不证明ERP一定没有对应商品")
        notes.extend(pf.get("qty_issues", []))
        result["核验说明"] = "；".join(notes)
    if code and not current:
        result["核验状态"] = "异常（ERP未找到）"
        result["应改为编码"] = ""
        result["匹配依据"] = "平台当前商家编码在本次ERP合并库中未找到；" + result["匹配依据"]
        result["核验说明"] = "当前编码未在ERP中找到；不能据此区分错码、旧码或导出范围问题。" + result["核验说明"]
    return result


def audit_components(text: str) -> dict[str, dict[str, Any]]:
    """Only compare named components, never treat every plus sign as a bundle."""
    return deepcopy(_cached_audit_components(text, parsing_rule_signature()))


@lru_cache(maxsize=32768)
def _cached_audit_components(text, rule_signature):
    text = quantity_prose(unicodedata.normalize('NFKC',text).lower())
    text = without_remark_gifts(text)
    text, _ = numeric_piece_offer(text,extract_variant_groups)
    text, _, _ = parenthetical_counts(text, COUNT_UNITS)
    parsed = quantity_structure(text)
    if parsed['additive']:
        components = {}
        kinds = Counter(p['key'].split('(')[0] for p in parsed['parts'])
        for part in parsed['parts']:
            base_kind = part['key'].split('(')[0]
            key = base_kind if kinds[base_kind] == 1 else part['key']
            feature = build_features('', part['text'])
            feature['families'] = set(part['families'])
            if feature['qty'] and feature['qty'][1] == '对' and feature['families'] & set(BUSINESS_RULES.get('pair_piece_families', [])):
                feature['qty'] = (feature['qty'][0] * 2, '片')
            if key in components and feature['qty'] and components[key]['qty'] and feature['qty'][1] == components[key]['qty'][1]:
                components[key] = {**components[key], 'qty': (components[key]['qty'][0] + feature['qty'][0], feature['qty'][1])}
            elif key not in components:
                components[key] = feature
        return components
    components = {}
    def same_product_gift(match):
        suffix = text[match.end():]
        if re.match(r"足贴|鞋垫|暖贴|眼罩|护膝|护腰", suffix):
            return match[0]
        return str(int(match[1]) + int(match[3])) + match[2]
    text = re.sub(r"(?:买)?(\d+)\s*(片|对|只|个)\s*(?:送|赠)\s*(\d+)\s*\2", same_product_gift, text)
    unit = re.search(r"(?:到手|共|合计|总计)\s*\d+\s*(片|对|只|个)", text)
    if unit:
        text = re.sub(r"买(\d+)送(\d+)(足贴|鞋垫|暖贴)", lambda m: m[1] + unit[1] + "+" + m[2] + unit[1] + m[3], text)
    text = re.sub(r"(?:到手|共|合计|总计)\s*\d+\s*(?:片|对|只|个)", "", text)
    for part in re.split(r"[+＋]|送|赠|搭配", text):
        kind = next((label for label, words in (
            ("鞋垫", ("鞋垫",)), ("足贴", ("足贴", "暖足贴")),
            ("眼罩", ("眼罩",)), ("护膝", ("护膝",)), ("护腰", ("护腰",)),
            ("暖膝", ("暖膝",)), ("暖腰", ("暖腰",)),
            ("暖手", ("暖手", "手机宝")), ("暖贴", ("暖贴", "暖宝宝")),
        ) if any(word in part for word in words)), "")
        if not kind or kind in components:
            continue
        features = build_features("", part)
        if features["qty"]:
            components[kind] = features
    return components


def audit_quantity(features: dict[str, Any]) -> tuple[int, str] | None:
    if 'confirmed_components' in features:
        return features['qty']
    text = features["quantity_text"]
    if re.search(r"(?:到手|共|合计|总计)\s*\d+", text):
        return features["qty"]
    # A per-box rate without the number of boxes does not establish a total.
    if re.search(r"\d+\s*(?:片|对|只|个)\s*[/／]\s*(?:盒|包|袋)", text) and features['qty'] is None:
        return None
    gift = re.search(r"(\d+)\s*(片|对|只|个)\s*(?:送|赠)\s*(\d+)\s*\2", text)
    if gift and len(_cached_audit_components(text, parsing_rule_signature())) <= 1:
        return int(gift[1]) + int(gift[3]), gift[2]
    return features["qty"]


def current_code_conflicts(pf: dict[str, Any], ef: dict[str, Any]) -> list[str]:
    """Find explicit contradictions; absent attributes are allowed by user policy."""
    ptext, etext = pf["quantity_text"], ef["quantity_text"]
    pe, ee = pf.get("evidence", {}), ef.get("evidence", {})
    conflicts = []
    # Thermal-patch use names overlap. Only clear cross-category differences are conflicts.
    family_group = lambda values: {"热贴用品" if value in {"暖贴", "足贴", "暖手用品", "暖手足宝"} else value for value in values}
    if pf["families"] and ef["families"] and not family_group(pf["families"]) & family_group(ef["families"]):
        conflicts.append("品类不一致：平台=" + "/".join(sorted(pf["families"])) + "，ERP=" + "/".join(sorted(ef["families"])))
    pparts, eparts = audit_feature_components(pf), audit_feature_components(ef)
    if len(pparts) >= 2 or len(eparts) >= 2:
        if len(pparts) >= 2 and len(eparts) >= 2 and set(pparts) != set(eparts):
            conflicts.append("套装组成品类不一致：平台=" + "+".join(pparts) + "，ERP=" + "+".join(eparts))
        for part in pparts.keys() & eparts.keys():
            if quantity_compare(pparts[part]["qty"], eparts[part]["qty"], pparts[part]["families"] | eparts[part]["families"]) is False:
                conflicts.append(f"套装{part}数量不一致：平台={pparts[part]['qty']}，ERP={eparts[part]['qty']}")
            for attr, label in [('formula','配方'),('scent','香型'),('material','材质'),('size','尺码')]:
                a = pparts[part].get('evidence', {}).get(attr, set())
                b = eparts[part].get('evidence', {}).get(attr, set())
                if attr == 'size':
                    normalize = lambda values: {re.sub(r'^x{2,4}l$',lambda m: str(len(m[0])-1)+'xl',x.lower()) for x in values}
                    a,b = normalize(a),normalize(b)
                if a and b and not a & b:
                    conflicts.append(f'套装{part}{label}不一致')
    elif feature_quantity_compare(pf,ef) is False:
        conflicts.append(f"数量不一致：平台={audit_quantity(pf)}，ERP={audit_quantity(ef)}")
    for key, label in (("formula", "配方/混合装"), ("material", "材质"),
                       ("size", "尺码"), ("gender", "男女童规格"),
                       ("capacity", "容量"), ("dimensions", "数字尺码/尺寸"), ("scent", "香型"),
                       ("model", "型号"), ("shape", "形状"), ("heat_type", "发热款型")):
        a, b = set(pe.get(key, set())), set(ee.get(key, set()))
        if key == "model" and not b:
            model = re.search(r"鞋套\s*(\d{3})(?!\d)", etext)
            b = {model.group(1)} if model else set()
        if key == "size":
            normalize = lambda values: {re.sub(r"^x{2,4}l$", lambda m: str(len(m[0])-1)+"xl", x.lower()) for x in values}
            a, b = normalize(a), normalize(b)
        if a and b and not a & b:
            conflicts.append(label + "不一致：平台=" + "/".join(sorted(a)) + "，ERP=" + "/".join(sorted(b)))
    # SKU color is selected data; title colors and unselected gender prose are not.
    color_values = lambda text: set(extract_variant_groups(text.replace("小黄鸭", "")).get("color", set())) | {x for x in ("咖", "肤", "橙") if x in text}
    a, b = color_values(pf.get('comparison_color_text',ptext)), color_values(ef.get('comparison_color_text',etext))
    if pf.get('business_color_values') is not None:
        a = set(pf['business_color_values'])
    if ef.get('business_color_values') is not None:
        b = set(ef['business_color_values'])
    if a and b and not a & b and not re.search(r"颜色随机|随机色|随机颜色", ptext):
        conflicts.append("颜色不一致：平台=" + "/".join(sorted(a)) + "，ERP=" + "/".join(sorted(b)))
    def body(text):
        text = text.replace('暖腕','手腕').replace('养护膝盖','膝盖')
        # A knee heating patch is not the same product type as a wearable
        # knee support. Only patch wording shares the thermal-patch label.
        text = re.sub(r'护膝(?:热敷)?贴|膝盖(?:热敷|保暖)贴','暖膝贴',text)
        parts={part for part in ('暖膝','暖腰','暖颈','手腕','护膝','护腰','暖足') if part in text}
        aliases=BUSINESS_RULES.get('body_part_aliases',{})
        return {next((canonical for canonical,words in aliases.items() if part in words),part) for part in parts}
    a, b = body(ptext) or body(pf["text"]), body(etext)
    if len(a) == len(b) == 1 and a != b:
        conflicts.append("主品部位不一致：平台=" + "/".join(a) + "，ERP=" + "/".join(b))
    def knee_type(feature):
        text=feature['quantity_text']
        if re.search(r'暖膝|膝盖.*(?:热敷|保暖).*贴|护膝(?:热敷)?贴',text):
            return '热敷贴'
        if '护膝' in text and (feature['evidence'].get('size') or re.search(r'\d+\s*(?:只|件|副)|穿戴|加绒',text)):
            return '穿戴护膝'
        return ''
    ptype,etype=knee_type(pf),knee_type(ef)
    if ptype and etype and ptype!=etype:
        conflicts.append('膝部商品类型不一致：穿戴护膝与暖膝热敷贴不能等同')
    for terms,label in [(('小魔仙','小黄鸭'),'暖芯系列'),(('大暖手','小暖手'),'暖芯大小款')]:
        a,b = {x for x in terms if x in ptext},{x for x in terms if x in etext}
        if a and b and not a & b:
            conflicts.append(label+'不一致')
    accessory_conflicts,_ = business_specs.accessory_conflicts(pf,ef)
    conflicts.extend(accessory_conflicts)
    return conflicts


def feature_quantity_compare(pf,ef):
    """Compare separately declared units directly, without granting conversion."""
    def values(f):
        result=set(tuple(q) for q in f.get('quantity_structure',{}).get('alternative_quantities',[]))
        qty=audit_quantity(f)
        if qty: result.add(tuple(qty))
        return result
    a,b=values(pf),values(ef)
    same=[x[0]==y[0] for x in a for y in b if x[1]==y[1]]
    if same: return all(same)
    return quantity_compare(audit_quantity(pf),audit_quantity(ef),pf['families']|ef['families'])


def composition_review_issues(pf, current):
    """Only explicit, unresolved composition is blocked; ordinary missing prose is unchanged."""
    problems = []
    for feature in [pf, *(r['features'] for r in current)]:
        detail = feature.get('quantity_structure', {})
        if detail.get('incomplete'):
            problems.extend(detail['issues'])
    for record in current:
        ef = record['features']
        if any(f.get('quantity_structure',{}).get('offer_details') for f in (pf,ef)):
            if feature_quantity_compare(pf,ef) is None:
                problems.append('买赠总片数与ERP数量尚未建立可比关系')
        if any(f.get('quantity_structure',{}).get('remark_gifts') for f in (pf,ef)):
            if feature_quantity_compare(pf,ef) is None:
                problems.append('精油仅作备注，主商品数量单位尚未建立可比关系')
        _,accessory_issues = business_specs.accessory_conflicts(pf,ef)
        problems.extend(accessory_issues)
        layered = any(f.get('quantity_structure', {}).get('packaging_details') for f in [pf,ef])
        if layered and len(audit_feature_components(pf)) <= 1 and len(audit_feature_components(ef)) <= 1:
            if quantity_compare(pf['qty'],ef['qty'],pf['families'] | ef['families']) is None:
                problems.append('外包装数量未建立可比对应关系')
        if pf['variants'].get('formula') == ef['variants'].get('formula') == {'三合一'} and pf['families'] == ef['families']:
            continue
        pparts, eparts = audit_feature_components(pf), audit_feature_components(ef)
        if len(pparts) >= 2 or len(eparts) >= 2:
            if set(pparts) != set(eparts):
                problems.append('平台与ERP组成尚未建立完整对应关系')
            for key in pparts.keys() & eparts.keys():
                a, b = pparts[key], eparts[key]
                if a['qty'] and b['qty'] and quantity_compare(a['qty'], b['qty'], a['families'] | b['families']) is None:
                    problems.append('组成单位不可比较：' + key)
    return list(dict.fromkeys(problems))


def packaging_review_issues(pf, current):
    issues = [r['features']['packaging_issue'] for r in current if r['features'].get('packaging_issue')]
    platform_issue = pf.get('packaging_issue')
    if platform_issue:
        explicit_family = extract_families(pf['quantity_text']) - {'包装材料'}
        resolved = any(r['features'].get('erp_product_confirmation') and explicit_family & r['features']['families']
            and quantity_compare(pf['qty'],r['features']['qty'],explicit_family|r['features']['families']) is True for r in current)
        if not resolved:
            issues.append(platform_issue)
    return list(dict.fromkeys(issues))


def audit_feature_components(feature):
    if 'confirmed_components' in feature:
        return feature['confirmed_components']
    # Only the outer mapping is extended below; cached features stay read-only.
    parts = dict(_cached_audit_components(feature.get('component_parse_text',feature['quantity_text']), parsing_rule_signature()))
    # Apparel exports often name the main product only in the title, while the
    # SKU specifies its size and bundled patches. Do not invent a piece count.
    primary = re.split('搭配|送|赠',feature['quantity_text'],maxsplit=1)[0]
    roles = [r for r in ('护膝', '护腰') if r in primary]
    if not roles:
        roles = [r for r in ('护膝', '护腰') if r in feature.get('title_text', '')]
    sizes = feature.get('variants', {}).get('size', set())
    if len(roles) == 1 and sizes and parts and re.search('搭配|送|赠',feature['quantity_text']):
        if roles[0] not in parts:
            parts[roles[0]] = build_features('', roles[0] + '/'.join(sorted(sizes)))
    return parts


def apply_confirmed_audit_specs(pf: dict[str, Any]) -> dict[str, Any]:
    """Use only already-confirmed product-scoped size/model relationships."""
    result = {**pf, "evidence": {k: set(values) for k, values in pf.get("evidence", {}).items()}}
    context, text = pf.get("context", {}), pf["quantity_text"]
    for rule in BUSINESS_RULES.get("confirmed_size_mappings", []):
        if rule.get("平台") != context.get("平台") or rule.get("商品 ID") != context.get("商品 ID"):
            continue
        scope = rule.get("平台款型", "")
        if any(word in scope and word not in text for word in ("中筒", "高筒")):
            continue
        material = norm_text(rule.get("材质", ""))
        if material and material not in result["evidence"].get("material", set()):
            continue
        selected = re.search(r"尺码[:：]\s*([^;,，]+)", text)
        value = rule["尺码对应"].get(selected.group(1).strip()) if selected else None
        if value:
            result["evidence"]["dimensions"] = {re.sub(r"[—–~／/]", "-", value)}
    for rule in BUSINESS_RULES.get("confirmed_product_models", []):
        if rule.get("平台") != context.get("平台") or rule.get("商品 ID") != context.get("商品 ID"):
            continue
        scope = rule.get("平台款型", "")
        if any(word in scope and word not in text for word in ("中筒", "高筒", "透明")):
            continue
        if norm_text(rule.get("材质", "")) not in result["evidence"].get("material", set()):
            continue
        result["evidence"]["model"] = {rule["ERP型号"]}
    return result


def audit_current_code(current_code: str, pf: dict[str, Any], erp_by_code: dict[str, list[dict[str, Any]]], erp_records: list[dict[str, Any]]) -> dict[str, Any]:
    code = literal_text(current_code).strip()
    prepared, issue = business_specs.transform_platform(pf, code, erp_by_code.get(code, []), build_features)
    result = _audit_current_code(current_code, prepared, erp_by_code, erp_records)
    if issue and result['核验状态'] == '核对通过':
        result.update(核验状态='待人工确认',核验说明=issue,冲突项目=issue)
    if prepared.get('business_rule_ids'):
        result['业务口径依据'] = '；'.join(prepared['business_rule_ids'])
        result['业务数量解析JSON'] = json.dumps(prepared['quantity_structure'], ensure_ascii=False)
    note = business_specs.pending_note(pf.get('context', {}))
    if note and code:
        result.update(核验状态='待人工确认',核验说明=note,匹配依据='用户明确保留待确认')
    agricultural = agriculture.audit(pf, literal_text(current_code).strip(), erp_by_code)
    if agricultural:
        result.update(agricultural)
    current = erp_by_code.get(literal_text(current_code).strip(), [])
    result = apply_weight_guard(result, pf, current)
    result = apply_cross_category_guard(result, prepared, current_code, current)
    return apply_quantity_parse_guard(result, pf, current_code, current)


def apply_cross_category_guard(result, pf, current_code, current):
    """Review the demonstrated thermal-piece / capacity-water-bottle ambiguity.

    This does not fill unknown families from SEO titles, grant unit conversions,
    or turn every missing attribute into a blocker.
    """
    if result['核验状态'] != '核对通过' or not current:
        return result
    confirmed = pf.get('confirmation')
    if confirmed and confirmed.get('record', {}).get('code') == literal_text(current_code).strip():
        return result
    specific = unicodedata.normalize('NFKC', pf['quantity_text'])
    quantity = audit_quantity(pf)
    if (not quantity or quantity[0] <= 0 or quantity[1] != '片'
            or pf.get('quantity_structure', {}).get('incomplete')):
        return result
    if '暖手' not in specific or not any(word in specific for word in ('穿戴', '草本', '暖手贴', '暖手宝贴', '蒸汽')):
        return result
    # A named water bottle/cover or multiple component types can be a real set.
    if any(word in specific for word in ('热水袋', '注水', '灌水', '布套', '袋套')):
        return result
    if len(audit_feature_components(pf)) >= 2:
        return result
    evidence = []
    for record in current:
        feature = record['features']
        if feature['families'] != {'热水袋'} or not feature.get('evidence', {}).get('capacity'):
            return result
        if len(audit_feature_components(feature)) >= 2:
            return result
        evidence.append({'code':record['code'], 'name':record['name'], 'source':record['source'],
                         'source_row':record.get('source_row'),
                         'capacity':sorted(feature['evidence']['capacity'])})
    note = (f'跨品类疑点，待人工确认：平台SKU明示暖手热贴类线索及{quantity[0]}片；'
            'ERP为有容量规格的热水袋：' + '；'.join(e['name'] for e in evidence)
            + '；未建立同一商品或合法主赠组成关系，未修改原编码')
    updated = dict(result)
    updated.update(核验状态='待人工确认', 核验说明=note, 匹配依据='SKU明示暖手片数与ERP热水袋容量品类存在疑点',
                   跨品类疑点JSON=json.dumps({'rule':'thermal_piece_vs_water_bottle_v1',
                       'platform_sku':pf['quantity_text'], 'platform_quantity':quantity, 'erp':evidence,
                       'original_algorithm_status':result['核验状态'],
                       'original_algorithm_reason':result.get('核验说明', '')}, ensure_ascii=False))
    updated['判定路径'] = updated.get('判定路径', '') + '；解析待确认：受限跨品类疑点'
    return updated


@lru_cache(maxsize=1)
def load_quantity_parse_guards():
    path = Path(__file__).resolve().parents[1] / 'config/quantity_parse_guards.json'
    data = json.loads(path.read_text(encoding='utf-8'))
    if data.get('version') != 1 or not isinstance(data.get('guards'), dict):
        raise ValueError('数量解析灰度配置必须为version=1且guards为对象')
    for key, entry in data['guards'].items():
        if (key != sku_reviews.review_key(entry['context'])
                or entry['content_fingerprint'] != sku_reviews.content_fingerprint(entry['context'])
                or not entry.get('code') or not entry.get('erp_fingerprints')):
            raise ValueError('数量解析灰度配置的身份或指纹无效')
    return data['guards']


def apply_quantity_parse_guard(result, pf, current_code, current):
    """Guard approved parsing gaps and expose expired bindings at the same ID."""
    context = pf.get('context', {})
    entry = load_quantity_parse_guards().get(sku_reviews.review_key(context))
    if not entry:
        return result
    reasons = []
    if entry['code'] != literal_text(current_code):
        reasons.append('当前编码变化，需重新确认')
    if entry['content_fingerprint'] != sku_reviews.content_fingerprint(context):
        changed = [k for k in sku_reviews.CONTENT if str(entry['context'].get(k, '')) != str(context.get(k, ''))]
        reasons.append('平台内容变化（' + '、'.join(changed) + '），需重新确认')
    if sorted(r.get('erp_fingerprint', '') for r in current) != entry['erp_fingerprints']:
        reasons.append('关联ERP记录缺失、重复或内容变化，需重新确认')
    if reasons:
        return mark_expired(result, '数量保护', '；'.join(reasons))
    if result['核验状态'] != '核对通过':
        return result
    comparison, _ = business_specs.transform_platform(pf, literal_text(current_code), current, build_features)
    comparison = apply_confirmed_audit_specs(comparison)

    def quantities(feature):
        found = {tuple(q) for q in feature.get('quantity_structure', {}).get('alternative_quantities', [])}
        quantity = audit_quantity(feature)
        if quantity:
            found.add(tuple(quantity))
        return found

    if not quantities(comparison):
        return result
    platform_parts = audit_feature_components(comparison)
    numeric_parts = {k for k, part in platform_parts.items() if quantities(part)}
    evidence = []
    for record in current:
        feature = record['features']
        # Missing declarations, incompatible units and already compared totals
        # are outside this rollout. Components can cover a missing aggregate.
        if quantities(feature) or feature_quantity_compare(comparison, feature) is not None:
            return result
        if not re.search(r'\d+\s*(?:盒|片|袋|包)', feature['quantity_text']):
            return result
        erp_parts = audit_feature_components(feature)
        if numeric_parts and numeric_parts.issubset(erp_parts) and all(
                feature_quantity_compare(platform_parts[k], erp_parts[k]) is not None for k in numeric_parts):
            return result
        issues = list(dict.fromkeys(feature.get('qty_issues', []) + feature.get('quantity_structure', {}).get('issues', [])))
        evidence.append({'code':record['code'], 'name':record['name'], 'source':record['source'],
                         'source_row':record.get('source_row'), 'issues':issues})
    if not evidence:
        return result
    note = 'ERP侧数量无法解析，未完成数量比对；' + '；'.join(
        e['name'] + '：' + '、'.join(e['issues'] or ['未获得可比较总量']) for e in evidence)
    result = dict(result)
    result.update(核验状态='待人工确认', 核验说明=note, 匹配依据='已确认16条数量解析灰度保护',
                  数量解析保护JSON=json.dumps({'platform_quantity':sorted(quantities(comparison)), 'erp':evidence}, ensure_ascii=False))
    result['判定路径'] = result.get('判定路径', '') + '；解析待确认：数量解析灰度保护'
    return result


def _audit_current_code(current_code: str, pf: dict[str, Any], erp_by_code: dict[str, list[dict[str, Any]]], erp_records: list[dict[str, Any]]) -> dict[str, Any]:
    code = literal_text(current_code).strip()
    if not code:
        result = classify_row(current_code, pf, erp_by_code, erp_records)
        result.update({"核验状态": "编码缺失（待补录）", "应改为编码": "", "冲突项目": "",
                       "匹配依据": "平台商家编码为空；" + result["匹配依据"],
                       "核验说明": "保留空编码记录；已有明确对应编码时另列补录参考，未明确时不猜测。"})
        return result
    current = erp_by_code.get(code, [])
    approved = pf.get("confirmation")
    result = {"核验状态": "异常（ERP未找到）", "应改为编码": "", "匹配依据": "平台原编码在本次ERP合并库中未找到",
              "核验说明": "不能据此断定是旧码、错码或ERP导出范围问题。", "ERP数据来源": " | ".join(dict.fromkeys(r["source"] for r in current)),
              "ERP唯一对应编码": "", "候选总数": 0, "独立身份命中条数": 0, "同规格候选数": 0,
              "当前编码ERP名称": " | ".join(r["name"] for r in current),
              "当前编码ERP来源": " | ".join(r["source"] for r in current), "当前编码ERP记录数": len(current),
              "数量解析": str(pf["qty"] or "未明确"), "规格问题": "；".join(pf.get("qty_issues", [])),
              "判定路径": "原编码收录与明示规格核对", "结构化可能ERP数": len(current), "结构化核验说明": "",
              "平台结构化属性": json.dumps({k: sorted(values) for k, values in pf.get("evidence", {}).items()}, ensure_ascii=False),
              "对应ERP结构化属性": "", "人工确认人": approved["confirmed_by"] if approved else "",
              "人工确认依据": approved["basis"] if approved else "", "冲突项目": ""}
    for i in range(3):
        suffix = "" if i == 0 else str(i+1)
        for field in ("名称", "编码", "来源", "分数", "冲突"):
            result[f"ERP候选{suffix}{field}"] = ""
    if not current:
        return result
    comparison = apply_confirmed_audit_specs(pf)
    comparisons = [current_code_conflicts(comparison, r["features"]) for r in current]
    if len(current) > 1:
        for i, a in enumerate(current):
            if any(current_code_conflicts(a["features"], b["features"]) for b in current[i+1:]):
                result.update({"核验状态": "待人工确认", "匹配依据": "ERP同一编码存在相互冲突的规格记录",
                               "核验说明": "保留全部ERP来源记录，不选取其中一条作为正确记录。", "冲突项目": "ERP同码规格冲突"})
                return result
    conflicts = list(dict.fromkeys(message for items in comparisons for message in items))
    result["对应ERP结构化属性"] = json.dumps([{k: sorted(values) for k, values in r["features"].get("evidence", {}).items()} for r in current], ensure_ascii=False)
    if conflicts:
        result.update({"核验状态": "异常（商品/规格不一致）", "匹配依据": "原编码已收录，但平台与关联ERP存在明确冲突",
                       "核验说明": "；".join(conflicts), "冲突项目": "；".join(conflicts)})
    else:
        result.update({"核验状态": "核对通过", "匹配依据": "原编码已收录，未发现明示商品或规格冲突",
                       "核验说明": "按用户口径：描述或规格信息不足也核对通过，不要求名称完全一致或全库身份唯一。"})
        composition_issues = composition_review_issues(comparison, current)
        if composition_issues:
            result.update({'核验状态':'待人工确认','匹配依据':'明确组合或包装表达未完成可比解析',
                           '核验说明':'需人工确认组成：'+'；'.join(composition_issues), '冲突项目':'；'.join(composition_issues)})
    if any(f.get('quantity_structure', {}).get('incomplete') for f in [comparison, *(r['features'] for r in current)]):
        reasons = composition_review_issues(comparison, current)
        result.update({'核验状态':'待人工确认','匹配依据':'明确组合或包装表达未完成可比解析',
                       '核验说明':'需人工确认组成：'+'；'.join(reasons), '冲突项目':'；'.join(dict.fromkeys([*conflicts,*reasons]))})
    packaging_issues = packaging_review_issues(comparison,current)
    if packaging_issues:
        result.update({'核验状态':'待人工确认','匹配依据':'成品或包装材料身份存在歧义',
                       '核验说明':'需人工确认成品/包材：'+'；'.join(packaging_issues),
                       '冲突项目':'；'.join(dict.fromkeys([*conflicts,*packaging_issues]))})
    if result['核验状态'] == '待人工确认' and (packaging_issues or composition_review_issues(comparison,current)):
        result['判定路径'] += '；解析待确认'
    result['数量解析备注'] = result['规格问题']
    result['规格问题'] = result['冲突项目']
    if pf.get("duplicate_sku"):
        result["核验说明"] += "；平台SKU ID重复，逐行保留并核对。"
    return result


def suggest_code_candidates(pf: dict[str, Any], result: dict[str, Any], records: list[dict[str, Any]], by_code: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    """Propose real ERP codes without changing the original audit or confirmations."""
    if agriculture.is_agricultural(pf) or result.get('农产品解析JSON'):
        return agriculture.suggestions(pf, result, records, by_code)
    output = {key: "" for key in ("推测商家编码", "推测ERP商品名称", "推测依据", "仍需确认的问题",
                                  "其他候选编码", "其他候选ERP记录", "推测ERP来源")}
    output['候选明细JSON'] = '[]'
    output["推测候选编码数"] = 0
    if result["核验状态"] not in {"异常（ERP未找到）", "异常（商品/规格不一致）", "编码缺失（待补录）", "待人工确认"}:
        return output
    comparison = apply_confirmed_audit_specs(pf)
    context, sku = pf.get("context", {}), pf["quantity_text"]
    def kind(text):
        return next((label for label, terms in (
            ("鞋套", ("鞋套",)), ("鞋垫", ("鞋垫",)), ("手套", ("手套",)),
            ("护膝", ("护膝",)), ("护腰", ("护腰",)), ("艾灸贴", ("艾灸", "艾热灸")),
            ("暖膝", ("暖膝",)), ("暖腰", ("暖腰",)), ("暖颈", ("暖颈",)), ("暖腕", ("暖腕", "手腕")),
            ("手机贴", ("手机宝", "手机加热贴", "手机暖贴")),
            ("暖手足宝", ("暖手足宝", "暖手宝", "暖手贴")),
            ("足贴", ("足贴", "暖足")), ("冰凉贴", ("冰凉贴",)), ("眼罩", ("眼罩",)),
            ("热水袋", ("热水袋",)), ("除湿", ("除湿", "干燥剂", "吸湿")),
            ("暖贴", ("暖贴", "暖宝宝", "热敷贴", "暖身贴")),
        ) if any(term in text for term in terms)), "")
    main_sku = re.split(r"[+＋]|搭配|送|赠", sku, maxsplit=1)[0]
    role = kind(main_sku) or kind(context.get("商品标题", pf["text"])) or kind(sku)
    def supported(record, confirmed=False):
        ef = record["features"]
        if not record["code"] or len(by_code.get(record["code"], [])) != 1:
            return False
        if not confirmed and (not role or role != kind(record['name']) or '包装材料' in ef['families']):
            return False
        if not confirmed:
            formulas = comparison.get('variants', {}).get('formula', set())
            if len(formulas)>1 and not formulas.issubset(ef.get('variants',{}).get('formula',set())):
                return False  # A mixed-formula SKU is not a single-formula candidate.
        if current_code_conflicts(comparison, ef):
            return False
        if packaging_review_issues(comparison, [record]):
            return False
        if not confirmed and role not in {'护膝', '护腰'} and composition_review_issues(comparison, [record]):
            return False
        if confirmed:
            return True
        if re.search(r"[+＋]|送|赠", record["name"]) and role not in {"护膝", "护腰"}:
            # Do not suggest a gift or mixed-material bundle for a plain SKU.
            pparts, eparts = audit_components(sku), audit_components(record["name"])
            if len(pparts) < 2 or set(pparts) != set(eparts):
                return False
        pe, ee = comparison["evidence"], ef["evidence"]
        # A replacement needs supporting specifications; absence is not supporting evidence.
        if role == "鞋套" and pe.get("size") and not pe.get("dimensions"):
            return False
        for key in ("formula", "material", "size", "heat_type", "scent"):
            if pe.get(key) and not ee.get(key):
                return False
        gender = ee.get('gender') or ef['variants'].get('gender', set())
        if pe.get('gender') and not pe['gender'] & gender:
            return False
        for key in ("model", "dimensions"):
            if role == "鞋套" and pe.get(key) and not ee.get(key):
                return False
        for term in ("经典", "增强", "高温", "国潮", "老人", "生姜"):
            if term in sku and term not in record["name"]:
                return False
        selected_colors = {c for c in ("咖", "肤", "蓝", "粉", "灰", "黑", "白", "透明", "橙", "黄", "红") if c in sku}
        if selected_colors and not re.search(r"颜色随机|随机色|随机颜色", sku):
            if not any(c in record["name"] for c in selected_colors):
                return False
        if role in {"护膝", "护腰"}:
            if not pe.get("size") or not ee.get("size"):
                return False
            pa, ea = audit_components(sku), audit_components(record["name"])
            return bool(pa.keys() & ea.keys()) and any(quantity_compare(pa[k]["qty"], ea[k]["qty"], pa[k]["families"] | ea[k]["families"]) is True for k in pa.keys() & ea.keys())
        if role == "鞋套":
            return bool(pe.get("dimensions") and ee.get("dimensions") and pe["dimensions"] & ee["dimensions"])
        return quantity_compare(audit_quantity(comparison), audit_quantity(ef), comparison["families"] | ef["families"]) is True

    candidates, basis = [], ""
    known_target = result.get("ERP唯一对应编码", "")
    if known_target and len(by_code.get(known_target, [])) == 1 and supported(by_code[known_target][0], confirmed=True):
        candidates, basis = [by_code[known_target][0]], "原核验已有明确补录对应编码，已查询本次ERP记录"
    approved = pf.get("confirmation")
    if not candidates and approved and supported(approved["record"], confirmed=True):
        candidates, basis = [approved["record"]], "有效已确认SKU关系，当前ERP记录指纹一致"
    if not candidates:
        for rule in BUSINESS_RULES.get("confirmed_style_bindings", []):
            if rule.get("平台") != context.get("平台") or rule["平台SKU"] != sku:
                continue
            if rule.get("商品 ID") and rule["商品 ID"] != context.get("商品 ID"):
                continue
            if rule.get("商品标题包含") and rule["商品标题包含"].lower() not in context.get("商品标题", "").lower():
                continue
            fingerprints = rule.get("允许ERP记录") or {rule["ERP编码"]:rule["ERP记录指纹"]}
            matches = [by_code[code][0] for code, fp in fingerprints.items() if len(by_code.get(code, [])) == 1 and by_code[code][0].get("erp_fingerprint") == fp]
            if len(matches) == len(fingerprints):
                candidates.extend(r for r in matches if supported(r, confirmed=True))
        if candidates:
            basis = "适用范围一致的已确认款式关系，ERP记录指纹有效"
    if not candidates:
        for rule in BUSINESS_RULES.get("confirmed_code_sets", []):
            if norm_text(rule["平台款型"]) not in norm_text(sku):
                continue
            if quantity_compare(audit_quantity(comparison), quantity_details(rule["平台总数量"])[0], comparison["families"]) is not True:
                continue
            for code in rule["允许编码"]:
                if len(by_code.get(code, [])) == 1 and supported(by_code[code][0], confirmed=True):
                    candidates.append(by_code[code][0])
        if candidates:
            basis = "已确认款型和总数量的允许编码范围，已查询本次ERP记录"
    if not candidates:
        candidates = [r for r in records if supported(r)]
        basis = "全库筛查主商品及明示规格，排除冲突后保留数量或尺码有支持的ERP记录"
    def evidence(record):
        ef = record['features']
        facts = []
        if role and role == kind(record['name']):
            facts.append('主商品类型一致')
        if quantity_compare(audit_quantity(comparison),audit_quantity(ef),comparison['families']|ef['families']) is True:
            facts.append('总数量一致')
        for key,label in [('formula','配方'),('material','材质'),('size','尺码'),('gender','男女童规格'),
                          ('capacity','容量'),('dimensions','数字尺寸'),('scent','香型'),('shape','形状'),
                          ('heat_type','发热款型'),('series','系列'),('duration','时长'),('brand','品牌')]:
            a = comparison['evidence'].get(key,set())
            b = ef['evidence'].get(key,set()) or (ef['variants'].get('gender',set()) if key=='gender' else set())
            if a and b and a & b:
                facts.append(label+'有相同明示值')
        colors = {c for c in ('咖','肤','蓝','粉','灰','黑','白','透明','橙','黄','红') if c in sku}
        if colors and not re.search(r'颜色随机|随机色|随机颜色',sku) and any(c in record['name'] for c in colors):
            facts.append('颜色有相同明示值')
        return facts
    candidates = list({r['code']:r for r in candidates}.values())
    proofs = {r['code']:evidence(r) for r in candidates}
    candidates.sort(key=lambda r:(-len(proofs[r['code']]),r['code']))
    output['候选明细JSON'] = json.dumps([{'code':r['code'],'name':r['name'],'source':r['source'],
        'source_row':r.get('source_row',''),'evidence':proofs[r['code']]} for r in candidates],ensure_ascii=False)
    output["推测候选编码数"] = len(candidates)
    if not candidates:
        output.update({"推测依据": "未找到满足当前品类及可比数量/尺码、且无明示冲突的可靠ERP候选",
                       "仍需确认的问题": "补充商品款式、规格或ERP记录；不凭名称相似度或旧编码格式猜码"})
        if role == '鞋套' and comparison['evidence'].get('size') and not comparison['evidence'].get('dimensions'):
            output['仍需确认的问题'] = '需明确该型号字母尺码与ERP数字范围的对应关系，未套用其他型号尺码表'
    elif len(candidates) == 1:
        chosen = candidates[0]
        output.update({"推测商家编码":chosen["code"], "推测ERP商品名称":chosen["name"],
                       "推测ERP来源":f"{chosen['source']}，ERP源行{chosen.get('source_row', '')}",
                       "推测依据":basis, "仍需确认的问题":"推测仅供审核，原编码和原异常状态保留，未保存为新的确认映射"})
    else:
        output.update({"其他候选编码":" / ".join(r["code"] for r in candidates),
                       "其他候选ERP记录":"\n".join(f"{r['code']}｜{r['name']}｜{r['source']}｜ERP源行{r.get('source_row','')}" for r in candidates),
                       "推测依据":basis, "仍需确认的问题":f"存在{len(candidates)}个合理候选，主推测编码留空，需选择实际款式、包装或编码"})
    return output


def infer_platform(filename: str) -> str:
    name = Path(filename).stem
    for token, platform in (("快手", "快手"), ("天猫", "天猫"), ("淘宝", "淘宝"), ("拼多多", "拼多多"), ("唯品会", "唯品会"), ("京东", "京东"), ("抖音", "抖音"), ("淘工厂", "淘工厂")):
        if token in name:
            return platform
    return "未识别平台"


def infer_store(filename: str) -> str:
    stem = Path(filename).stem
    if "快手小店商品导出" in stem:
        return "快手小店"
    # Common export timestamp suffixes are removed only for display; source filename is retained separately.
    stem = re.sub(r"[-_ ]?\d{4}[-_]?\d{2}[-_]?\d{2}.*$", "", stem)
    stem = re.sub(r"商品导出$", "", stem)
    return stem.strip("-_ ") or Path(filename).stem


def load_platform(path: str, explicit_sheet: str | None = None) -> tuple[pd.DataFrame, str, int]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        df, header_row = detect_csv_header(path)
        return df.fillna(""), "CSV", header_row
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        df, sheet, header_row = detect_excel_sheet(path, explicit_sheet=explicit_sheet)
        return df.fillna(""), sheet, header_row
    raise RuntimeError(f"暂不支持的文件格式：{suffix}；请使用 XLSX、XLS 或 CSV。")


def build_erp_records(erp_df: pd.DataFrame) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    records = []
    by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for _, row in erp_df.iterrows():
        code = str(row.get("ERP商家编码", "")).strip()
        name = str(row.get("ERP商品名称", "")).strip()
        category = str(row.get("ERP类别", "")) + " " + str(row.get("ERP分类", ""))
        features = build_features("", "", erp_name=name, erp_category=category)
        explicit_quantity = str(row.get("ERP规格数量", ""))
        if explicit_quantity:
            features["qty"], features["qty_issues"] = quantity_details(explicit_quantity)
        record = {"code": code, "name": name, "source": str(row.get("ERP来源", "")),
                  "item_id": str(row.get("ERP货品编号", "")).strip(),
                  "spec_code": str(row.get("ERP规格编码", "")).strip(),
                  "source_row": int(row["原始行号"]), "features": features}
        record['net_weight_kg'] = str(row.get('ERP重量', '')).strip()
        record["erp_fingerprint"] = confirmations.digest({"source": record["source"], "raw": json.loads(row["原始字段JSON"])})
        records.append(record)
        if code:
            by_code[code].append(record)
    interpretation_path = Path(__file__).resolve().parents[1] / 'config/erp_product_confirmations.json'
    interpretations = json.loads(interpretation_path.read_text(encoding='utf-8'))['records'] if interpretation_path.exists() else {}
    for code, interpretation in interpretations.items():
        matches = by_code.get(code, [])
        if len(matches) == 1 and matches[0]['erp_fingerprint'] == interpretation['erp_fingerprint']:
            matches[0]['features'] = {**matches[0]['features'], 'families':{interpretation['family']},
                'packaging_issue':'', 'erp_product_confirmation':interpretation['basis']}
    for record in records:
        business_specs.transform_erp(record, build_features)
    return records, by_code


def verify_platform_file(path: str, erp_records: list[dict[str, Any]], erp_by_code: dict[str, list[dict[str, Any]]], store_override: str | None = None, sheet_override: str | None = None, platform_override: str | None = None, confirmed_entries: dict | None = None, review_entries: dict | None = None) -> tuple[pd.DataFrame, dict[str, Any]]:
    platform = platform_override or infer_platform(path)
    profile = PLATFORM_PROFILES.get(platform, {})
    df, sheet, header_row = load_platform(path, explicit_sheet=sheet_override or profile.get("sheet"))
    cols = choose_platform_columns(df, platform)
    store = store_override or infer_store(path)
    confirmed_entries = confirmations.load_mappings() if confirmed_entries is None else confirmed_entries
    review_entries = sku_reviews.load_reviews() if review_entries is None else review_entries
    sku_column = cols.get("sku_id")
    duplicate_ids = set()
    if sku_column:
        ids = df[sku_column].map(literal_text).str.strip()
        duplicate_ids = set(ids[ids.ne("") & ids.duplicated(keep=False)])
    out_rows, raw_rows, quantity_rows = [], [], []
    for idx, row in df.iterrows():
        if not any(str(value).strip() for value in row.tolist()):
            continue
        def value(field):
            if field == 'sku_name' and profile.get('sku_name_columns'):
                missing = [c for c in profile['sku_name_columns'] if c not in df.columns]
                if missing:
                    raise RuntimeError(f"平台{platform}的SKU属性列缺失：{missing}；请确认导出模板。")
                return '；'.join(str(row[c]) for c in profile['sku_name_columns'] if str(row[c]).strip())
            return str(row[cols[field]]) if cols.get(field) else ""
        product_id, title, sku_id, sku_name, current_code = (value(k) for k in ("product_id", "title", "sku_id", "sku_name", "code"))
        pf = build_features(title, sku_name)
        if platform == '京东' and '店内分类' in df.columns:
            pf['agricultural'] = agriculture.parse_spec(sku_name, title, str(row['店内分类']))
        pf["erp_item_id"] = literal_text(value("erp_item_id")).strip()
        pf["spec_code"] = literal_text(value("spec_code")).strip()
        pf["duplicate_sku"] = literal_text(sku_id).strip() in duplicate_ids
        if cols.get("quantity"):
            pf["qty"], pf["qty_issues"] = quantity_details(value("quantity"))
        context = {"平台": platform, "店铺": store, "商品 ID": product_id, "SKU ID": sku_id,
                   "商品标题": title, "SKU": sku_name, "ERP货品编号": pf["erp_item_id"],
                   "ERP规格编码": pf["spec_code"], "独立规格数量": value("quantity")}
        if platform == '京东' and '店内分类' in df.columns:
            context['农产品店内分类'] = str(row['店内分类'])
        pf["context"] = context
        component_rule = BUSINESS_RULES.get("component_confirmations", {}).get(confirmations.platform_key(context))
        if component_rule and component_rule["platform_fingerprint"] == confirmations.platform_fingerprint(context):
            pf["evidence"]["formula_composition"] = set(component_rule["composition"])
            pf["component_basis"] = component_rule["basis"]
        pf["confirmation"], pf["confirmation_issue"] = confirmations.lookup_confirmation({**context, "原编码": current_code}, erp_records, confirmed_entries)
        result = audit_current_code(current_code, pf, erp_by_code, erp_records)
        result, reviewed_suggestion = sku_reviews.apply_review(context, current_code, result, erp_by_code, review_entries)
        pending_note = business_specs.pending_note(context)
        if pending_note and literal_text(current_code).strip():
            result.update(核验状态='待人工确认',核验说明=pending_note,匹配依据='用户明确保留待确认')
            reviewed_suggestion = sku_reviews.empty_suggestion(pending_note)
        current_records = erp_by_code.get(literal_text(current_code).strip(), [])
        agricultural = agriculture.audit(pf, literal_text(current_code).strip(), erp_by_code)
        if agricultural:
            result.update(agricultural)
            reviewed_suggestion = None
        result = apply_weight_guard(result, pf, current_records)
        result = apply_quantity_parse_guard(result, pf, current_code, current_records)
        if (result.get('数量解析保护JSON') or result.get('确认失效JSON')
                or (result.get('跨品类疑点JSON') and result['核验状态'] == '待人工确认')):
            reviewed_suggestion = sku_reviews.empty_suggestion(result['核验说明'])
        quantity_rows.append({'商品 ID':product_id, 'SKU ID':sku_id, '源行号':idx + header_row + 2,
            'platform':pf['quantity_structure'],
            'confirmed_business_specs':json.loads(result['业务数量解析JSON']) if result.get('业务数量解析JSON') else None,
            'agricultural':json.loads(result['农产品解析JSON']) if result.get('农产品解析JSON') else None,
            'erp':[{'code':r['code'], 'source':r['source'], 'name':r['name'],
                    'quantity':r['features']['quantity_structure']} for r in current_records]})
        out_rows.append({"店铺": store, "平台": platform, "商品 ID": product_id, "SKU ID": sku_id,
                         "SKU": sku_name, "原编码": current_code, "应改为编码": result["应改为编码"],
                         "核验状态": result["核验状态"], "匹配依据": result["匹配依据"],
                         "ERP数据来源": result["ERP数据来源"], "异常说明": result["核验说明"],
                         "商品标题": title, "ERP货品编号": pf["erp_item_id"], "ERP规格编码": pf["spec_code"],
                         "独立规格数量": value("quantity"), "平台内容指纹": confirmations.platform_fingerprint(context),
                         "组成确认依据": pf.get("component_basis", ""),
                         **{k: val for k, val in result.items() if k not in {"应改为编码", "核验状态", "匹配依据", "ERP数据来源", "核验说明"}},
                         "源文件": str(Path(path).resolve()), "源工作表": sheet, "源行号": idx + header_row + 2})
        out_rows[-1].update(reviewed_suggestion if reviewed_suggestion is not None else suggest_code_candidates(pf, result, erp_records, erp_by_code))
        raw_rows.append({"来源文件": str(Path(path).resolve()), "来源工作表": sheet,
                         "来源行号": idx + header_row + 2, "原始字段JSON": json.dumps(row.to_dict(), ensure_ascii=False)})
    meta = {"store": store, "platform": platform, "source_file": str(Path(path).resolve()), "sheet": sheet,
            "header_row": header_row, "source_rows": len(out_rows), "columns": json.dumps(cols, ensure_ascii=False),
            "sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(), "_raw_rows": raw_rows,
            "_quantity_details": quantity_rows}
    return pd.DataFrame(out_rows), meta


def style_workbook(path: str) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path)
    status_fills = {"核对通过": "E2F0D9", "异常（商品/规格不一致）": "FFC7CE", "编码正确": "E2F0D9", "编码错误（建议修改）": "FCE4D6",
                    "编码缺失（待补录）": "FFF2CC", "待人工确认": "F4CCCC", "ERP无匹配": "E4DFEC", "异常（ERP未找到）": "FFC7CE"}
    long_fields = {"商品标题", "SKU", "当前编码ERP名称", "ERP候选名称", "ERP候选2名称", "ERP候选3名称",
                   "匹配依据", "异常说明", "核验说明", "说明", "组合装名称", "货品名称", "ERP商品名称", "对应ERP名称", "结构化核验说明", "人工确认依据", "确认依据", "组成确认依据", "平台结构化属性", "对应ERP结构化属性", "冲突项目", "数量解析备注", "推测ERP商品名称", "推测依据", "仍需确认的问题", "其他候选ERP记录", "其他候选编码"}
    for ws in wb:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "E2" if ws.title in {"核验结果", "待人工确认", "编码待补录", "异常（ERP未找到）", "异常（商品规格不一致）", "异常及补录推测"} else "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = {cell.column: str(cell.value or "") for cell in ws[1]}
        widths = {}
        for index, header in headers.items():
            if header == "原始字段JSON":
                width = 96
            elif header in long_fields:
                width = 64 if header not in {"说明", "异常说明"} else 80
            elif any(token in header.lower() for token in ("文件", "source_file", "sha256")):
                width = 68
            elif any(token in header.lower() for token in ("编码", "sku", "id", "货品编号")):
                width = 26
            elif header in {"店铺", "平台", "ERP来源", "ERP数据来源", "核验状态", "ERP候选来源"}:
                width = 24
            else:
                width = max(16, min(28, len(header) * 2 + 2))
            widths[index] = width
            ws.column_dimensions[get_column_letter(index)].width = width
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            lines = 1
            for cell in row:
                header = headers[cell.column]
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if any(token in header.lower() for token in ("编码", "sku", "id", "货品编号")):
                    cell.number_format = "@"
                cell.font = Font(name="Microsoft YaHei", size=10, color="222222")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if header == "核验状态" and cell.value in status_fills:
                    cell.fill = PatternFill("solid", fgColor=status_fills[cell.value])
                units = sum(2 if ord(c) > 255 else 1 for c in str(cell.value or ""))
                cell_lines = (units + int(widths[cell.column]) - 3) // max(1, int(widths[cell.column]) - 2)
                # The complete raw JSON is retained and inspectable in the
                # formula bar; keep this audit table usable for scrolling.
                if header == "原始字段JSON":
                    cell_lines = min(cell_lines, 6)
                lines = max(lines, cell_lines)
            ws.row_dimensions[row[0].row].height = min(409, max(28, lines * 14 + 6))
    wb.save(path)


def write_report(output_path: str, results: pd.DataFrame, erp_df: pd.DataFrame, metas: list[dict[str, Any]]) -> None:
    status_counts = results.groupby(["店铺", "平台", "核验状态"], dropna=False).size().reset_index(name="记录数") if len(results) else pd.DataFrame(columns=["店铺", "平台", "核验状态", "记录数"])
    total_counts = results.groupby(["店铺", "平台"], dropna=False).size().reset_index(name="商品SKU总数") if len(results) else pd.DataFrame(columns=["店铺", "平台", "商品SKU总数"])
    status_pivot = results.pivot_table(index=["店铺", "平台"], columns="核验状态", values="SKU", aggfunc="count", fill_value=0).reset_index() if len(results) else pd.DataFrame()
    suggestion_counts = []
    if '推测商家编码' in results:
        for (store, platform), rows in results.groupby(['店铺', '平台'], dropna=False):
            pending = rows[rows['核验状态'].isin(['异常（ERP未找到）', '异常（商品/规格不一致）', '编码缺失（待补录）', '待人工确认'])]
            suggestion_counts.append({'店铺':store, '平台':platform,
                '推测编码行数':int(pending['推测商家编码'].ne('').sum()),
                '多候选行数':int(pending['其他候选编码'].ne('').sum()),
                '无可靠候选行数':int((pending['推测候选编码数'].eq(0)).sum())})
    summary = total_counts.merge(status_pivot, on=['店铺','平台'])
    if suggestion_counts:
        summary = summary.merge(pd.DataFrame(suggestion_counts), on=['店铺','平台'])
    erp_code_counts = erp_df.groupby("ERP商家编码", dropna=False).size().reset_index(name="记录数")
    duplicate_codes = erp_code_counts[(erp_code_counts["ERP商家编码"].astype(str).str.strip() != "") & (erp_code_counts["记录数"] > 1)]
    blank_code_count = int(erp_df["ERP商家编码"].astype(str).str.strip().eq("").sum())
    erp_summary = pd.DataFrame([
        {"检查项": "组合装记录数", "结果": int((erp_df["ERP来源"] == "组合装").sum()), "说明": "完整纳入合并 ERP 库"},
        {"检查项": "单品记录数", "结果": int((erp_df["ERP来源"] == "单品").sum()), "说明": "完整纳入合并 ERP 库"},
        {"检查项": "ERP合并记录总数", "结果": int(len(erp_df)), "说明": "未按名称或编码擅自去重"},
        {"检查项": "ERP非空商家编码唯一值", "结果": int(erp_df.loc[erp_df["ERP商家编码"].astype(str).str.strip().ne(""), "ERP商家编码"].nunique()), "说明": "以原始字符串保留编码格式"},
        {"检查项": "ERP空商家编码记录", "结果": blank_code_count, "说明": "需检查是否为源数据问题"},
        {"检查项": "ERP重复商家编码组数", "结果": int(len(duplicate_codes)), "说明": "重复记录不删除；同码规格明确冲突时待人工确认，无冲突则按原编码核对"},
    ])
    duplicate_codes = duplicate_codes.merge(erp_df[["ERP商家编码", "ERP来源", "ERP商品名称"]], on="ERP商家编码", how="left") if len(duplicate_codes) else pd.DataFrame(columns=["ERP商家编码", "记录数", "ERP来源", "ERP商品名称"])
    rules = pd.DataFrame([
        {'规则': 'ERP来源', '说明': '同时读取组合装和单品CSV，保留全部行、原字段和来源，不按名称或编码去重。'},
        {'规则': '核对通过', '说明': '原编码已收录且无明示冲突时，普通描述信息不足可通过；明确组成解析失败、成品/包材歧义和农产品重量待核不放行。'},
        {'规则': '异常（ERP未找到）', '说明': '平台当前非空编码在ERP商家编码列未找到；不自动认定是旧码、错码或ERP导出范围问题。'},
        {'规则': '异常（商品/规格不一致）', '说明': '当前编码已收录，但双方有明确数量、品类、配方、尺码、颜色、材质、容量、型号、香型或套装组成冲突；冲突项目列写实际依据。'},
        {'规则': '信息不足', '说明': '普通属性缺失不当作明确冲突；明确组成或包装表达解析失败进入待确认；农产品重量注明需人工核重，不自动换算。'},
        {'规则': '数量与套装', '说明': '主品、赠品、总数分别处理，不将总数算到赠品上；足贴、发热鞋垫、暖手足宝按已确认规则1对=2片。'},
        {'规则': '三合一', '说明': '混合装之间只比较总数，忽略配比；混合装对单配方仍属明确冲突，用户逐行确认的7条试用装例外仅限原绑定SKU。'},
        {'规则': '空编码', '说明': '全部保留并列编码缺失（待补录）。已有明确关系可列ERP唯一对应编码，未明确时不猜测。'},
        {'规则': '同码与重复SKU', '说明': 'ERP同码规格明确冲突时待人工确认，同码同规格保留来源且按普通编码核对；平台重复SKU逐行保留。'},
        {'规则': '建议编码', '说明': '异常和空码追加推测编码、ERP名称、依据、仍需确认的问题及其他候选。单一有支持的候选才填主推测码，多候选主列留空，未知也留空；原编码、状态和应改为编码保持不变，不写入确认映射。'},
        {'规则': '字段与原值', '说明': '快手读取sku编码，天猫读取SKU分组商家编码.1；编码和ID按文本保存，保留前导零及原文，不执行公式。'},
        {'规则': '已确认关系', '说明': '保留原有映射用于追踪及明确空码的补录参考；没有映射或映射过期，不阻止已收录且无明示冲突的编码通过。'},
        {'规则': '准确率', '说明': '核对通过包含用户允许的信息不足情况；没有全量人工标准答案，不报告真实商品匹配准确率。'},
    ])
    meta_df = pd.DataFrame([{k: v for k, v in meta.items() if not k.startswith("_")} for meta in metas])
    review = results[results["核验状态"].isin(["待人工确认", "ERP无匹配"])].copy() if len(results) else results.copy()
    # Preserve all source columns from both ERP datasets in the merged library sheet.
    erp_out = erp_df.drop(columns=["_name_norm"], errors="ignore")
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="核验汇总", index=False)
        results.to_excel(writer, sheet_name="核验结果", index=False)
        if '推测商家编码' in results:
            suggestion_columns = ['店铺','平台','商品 ID','SKU ID','商品标题','SKU','原编码','核验状态',
                                  '当前编码ERP名称','异常说明','推测商家编码','推测ERP商品名称','推测依据',
                                  '仍需确认的问题','其他候选编码','其他候选ERP记录','推测ERP来源','推测候选编码数',
                                  '源文件','源工作表','源行号']
            results.loc[results['核验状态'].isin(['异常（ERP未找到）','异常（商品/规格不一致）','编码缺失（待补录）','待人工确认']), suggestion_columns].to_excel(writer, sheet_name='异常及补录推测', index=False)
            candidate_detail_columns = ['店铺','平台','商品 ID','SKU ID','SKU','原编码','核验状态','候选类型','候选ERP编码','候选ERP商品名称','ERP来源','ERP源行号','推测依据','仍需确认的问题']
            erp_lookup = {literal_text(row['ERP商家编码']).strip():row for row in erp_df.to_dict('records')}
            candidate_details = []
            for row in results.to_dict('records'):
                codes = [row['推测商家编码']] if row['推测商家编码'] else row['其他候选编码'].split(' / ') if row['其他候选编码'] else []
                for code in codes:
                    source = erp_lookup[code]
                    candidate_details.append({**{key:row[key] for key in ['店铺','平台','商品 ID','SKU ID','SKU','原编码','核验状态','推测依据','仍需确认的问题']},
                        '候选类型':'单一参考' if row['推测商家编码'] else '多个候选', '候选ERP编码':code,
                        '候选ERP商品名称':source['ERP商品名称'], 'ERP来源':source['ERP来源'], 'ERP源行号':source['原始行号']})
            pd.DataFrame(candidate_details, columns=candidate_detail_columns).to_excel(writer, sheet_name='推测候选明细', index=False)
        status_counts.to_excel(writer, sheet_name="状态汇总明细", index=False)
        review.to_excel(writer, sheet_name="待人工确认", index=False)
        supplement_columns = list(results.columns)
        supplement_columns.remove("ERP唯一对应编码")
        supplement_columns.insert(supplement_columns.index("原编码") + 1, "ERP唯一对应编码")
        results.loc[results["核验状态"].eq("编码缺失（待补录）"), supplement_columns].to_excel(writer, sheet_name="编码待补录", index=False)
        confirmation_columns = [*confirmations.CONTEXT_FIELDS, "平台内容指纹", "原编码", "核验状态", "源文件", "源工作表", "源行号", "ERP候选编码", "ERP候选名称"]
        form = review[confirmation_columns].copy()
        for column in ("确认ERP编码", "确认人", "确认依据", "确认"):
            form[column] = ""
        form.to_excel(writer, sheet_name="人工确认表", index=False)
        results[results["核验状态"].eq("异常（ERP未找到）")].to_excel(writer, sheet_name="异常（ERP未找到）", index=False)
        results[results["核验状态"].eq("异常（商品/规格不一致）")].to_excel(writer, sheet_name="异常（商品规格不一致）", index=False)
        erp_summary.to_excel(writer, sheet_name="ERP合并校验", index=False)
        duplicate_codes.to_excel(writer, sheet_name="ERP重复编码", index=False)
        rules.to_excel(writer, sheet_name="核验规则", index=False)
        foot_rules = BUSINESS_RULES.get("confirmed_code_sets", [])
        if foot_rules:
            pd.DataFrame([{**rule, "允许编码": " / ".join(rule["允许编码"])} for rule in foot_rules]).to_excel(writer, sheet_name="已确认足贴编码", index=False)
        meta_df.to_excel(writer, sheet_name="平台文件识别", index=False)
        erp_out.to_excel(writer, sheet_name="ERP合并库", index=False)
        pd.DataFrame([r for meta in metas for r in meta.get("_raw_rows", [])]).to_excel(writer, sheet_name="平台原始数据", index=False)
    style_workbook(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="平台原商家编码收录及明确商品规格冲突核验")
    parser.add_argument("--version", action="version", version=VERSION)
    parser.add_argument("--combo-erp", help="本次组合装CSV；与单品同时指定，都省略则使用固定标准库")
    parser.add_argument("--single-erp", help="本次单品CSV；与组合装同时指定，不修改固定标准库")
    parser.add_argument("--platform-file", action="append", default=[], help="平台商品导出文件；可重复指定多个文件")
    parser.add_argument("--store", action="append", default=None, help="店铺名称；可重复指定，顺序对应 --platform-file")
    parser.add_argument("--sheet", action="append", default=None, help="指定平台 Excel 工作表；可重复指定，顺序对应 --platform-file")
    parser.add_argument("--output", default="商品编码核验结果.xlsx", help="输出 XLSX 文件路径")
    parser.add_argument("--platform", action="append", help="明确平台名，与平台文件一一对应；省略则按文件名识别")
    parser.add_argument("--mapping-config", help="字段映射JSON路径；默认使用项目config/platform_aliases.json")
    parser.add_argument("--confirmed-mappings", help="人工确认映射JSON；默认config/confirmed_mappings.json")
    parser.add_argument("--import-confirmations", help="从已填写人工确认表的报告导入确认，需同时提供两份ERP；本模式不输入平台文件")
    args = parser.parse_args()
    try:
        args.combo_erp, args.single_erp, library_info = resolve_erp_inputs(args.combo_erp,args.single_erp)
    except (ValueError,OSError) as exc:
        parser.error(str(exc))
    if args.mapping_config:
        load_alias_config(args.mapping_config)
    mapping_path = Path(args.confirmed_mappings) if args.confirmed_mappings else confirmations.DEFAULT_PATH
    protected = [args.combo_erp, args.single_erp, *args.platform_file]
    if args.import_confirmations:
        protected.append(args.import_confirmations)
    if mapping_path.resolve() in {Path(p).resolve() for p in protected} | {Path(args.output).resolve()}:
        parser.error("人工映射文件不能覆盖输入或输出报告。")
    if args.import_confirmations:
        if args.platform_file:
            parser.error("导入人工确认时不同时核验平台文件。")
        erp_df = load_erp(args.combo_erp, args.single_erp)
        records, _ = build_erp_records(erp_df)
        count = confirmations.import_confirmations(args.import_confirmations, records, erp_df, mapping_path)
        print(f"已保存{count}条明确人工确认；映射文件：{mapping_path}")
        return 0
    if not args.platform_file:
        parser.error("核验模式至少需要一个--platform-file。")
    if args.platform and len(args.platform) != len(args.platform_file):
        parser.error("--platform数量必须与--platform-file数量一致。")
    if Path(args.output).resolve() in {Path(p).resolve() for p in [args.combo_erp, args.single_erp, *args.platform_file]}:
        parser.error("输出路径不能覆盖ERP或平台原始文件。")
    if Path(args.output).suffix.lower() != ".xlsx":
        parser.error("输出文件必须使用.xlsx扩展名。")
    audit_path = Path(args.output).parent/'核验依据'/(Path(args.output).stem+'.json')
    extra_protected = [mapping_path,DEFAULT_MANIFEST,Path(args.mapping_config)] if args.mapping_config else [mapping_path,DEFAULT_MANIFEST]
    if audit_path.resolve() in {Path(p).resolve() for p in [*protected,*extra_protected]}:
        parser.error('内部核验依据不能覆盖输入或配置文件。')
    Path(args.output).resolve().parent.mkdir(parents=True, exist_ok=True)
    if args.store and len(args.store) not in {1, len(args.platform_file)}:
        parser.error("--store 可指定一次（单文件）或按每个 --platform-file 的顺序各指定一次。")
    if args.sheet and len(args.sheet) not in {1, len(args.platform_file)}:
        parser.error("--sheet 可指定一次（单文件）或按每个 --platform-file 的顺序各指定一次。")
    if len(args.platform_file) > 1 and args.store and len(args.store) == 1:
        parser.error("批量处理多个文件时，请对每个 --platform-file 分别指定 --store，或全部省略。")
    if len(args.platform_file) > 1 and args.sheet and len(args.sheet) == 1:
        parser.error("批量处理多个文件时，请对每个 --platform-file 分别指定 --sheet，或全部省略。")
    erp_df = load_erp(args.combo_erp, args.single_erp)
    erp_records, erp_by_code = build_erp_records(erp_df)
    confirmed_entries = confirmations.load_mappings(mapping_path)
    all_results, metas = [], []
    for i, path in enumerate(args.platform_file):
        store_override = args.store[i] if args.store else None
        sheet_override = args.sheet[i] if args.sheet else None
        result, meta = verify_platform_file(path, erp_records, erp_by_code, store_override=store_override, sheet_override=sheet_override, platform_override=args.platform[i] if args.platform else None, confirmed_entries=confirmed_entries)
        all_results.append(result)
        metas.append(meta)
    # Result columns combine numeric counts and text; missing fields use an empty
    # string. Choose a compatible dtype before filling, without hiding warnings.
    results = pd.concat(all_results, ignore_index=True).astype(object) if all_results else pd.DataFrame()
    results = results.where(results.notna(), '')
    write_simple_report(args.output, results)
    audit_path.parent.mkdir(parents=True,exist_ok=True)
    audit_path.write_text(json.dumps({'version':VERSION,'erp_library':library_info,
        'erp_rows':len(erp_df),'platform_sources':metas,'results':results.to_dict('records')},ensure_ascii=False,indent=2),encoding='utf-8')
    print(f"输出完成：{args.output}")
    print(f"ERP 合并记录：{len(erp_df)}；平台 SKU 核验行：{len(results)}")
    if len(results):
        print(results.groupby(["店铺", "核验状态"]).size().to_string())
    return 0

if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise
