"""The default two-sheet delivery, independent of chat apps or Node.js."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

HEADERS = ['平台','商品ID','SKU ID','商品名称','SKU规格','原编码','核验结果','推测编码','问题说明','候选1','候选2','候选3']
STATUS = {'核对通过':'通过','异常（ERP未找到）':'ERP未找到','异常（商品/规格不一致）':'规格不一致',
          '编码缺失（待补录）':'编码为空','待人工确认':'待确认'}


def prepare_rows(results):
    rows = []
    for original in results.to_dict('records'):
        state = original['核验状态']
        candidates = json.loads(original.get('候选明细JSON') or '[]')
        multiple = len(candidates) > 1
        if state == '核对通过':
            note = ''
        elif state == '异常（商品/规格不一致）':
            note = original.get('冲突项目') or original.get('异常说明') or ''
            note = note.replace('颜色不一致：平台=咖，ERP=灰','平台咖色，ERP灰色')
            note = note.replace('配方/混合装不一致：平台=三合一，ERP=艾草','平台三合一，ERP单配方艾草')
        elif multiple:
            note = f'{len(candidates)}个候选，待确认' + ('（展示3个）' if len(candidates)>3 else '')
        elif original.get('推测商家编码'):
            note = '原编码未收录' if state == '异常（ERP未找到）' else '待补录'
        else:
            note = '暂无可靠推测，待确认'
        if any(marker in original.get('判定路径', '') for marker in ('人工SKU复核','解析待确认','需人工核重','确认失效')) and state != '核对通过':
            note = original.get('异常说明') or note
        row = {'平台':original.get('平台',''),'商品ID':original.get('商品 ID',''),'SKU ID':original.get('SKU ID',''),
            '商品名称':original.get('商品标题',''),'SKU规格':original.get('SKU',''),'原编码':original.get('原编码',''),
            '核验结果':STATUS.get(state,state),'推测编码':original.get('推测商家编码',''),'问题说明':note}
        for i in range(3):
            row[f'候选{i+1}'] = f"{candidates[i]['code']}\n{candidates[i]['name']}" if multiple and i<len(candidates) else ''
        rows.append(row)
    priority = {'规格不一致':0,'ERP未找到':1,'编码为空':2,'待确认':3}
    pending = [row for row in rows if row['核验结果'] != '通过']
    pending.sort(key=lambda row:(priority.get(row['核验结果'],4),row['问题说明'],row['SKU规格'],row['平台'],row['商品ID']))
    return pending, rows


def write_simple_report(output_path, results):
    pending, all_rows = prepare_rows(results)
    book = Workbook()
    book.remove(book.active)
    widths = [9,20,22,36,42,22,16,22,32,34,34,34]
    fills = {'规格不一致':'FCE4D6','ERP未找到':'FCE4D6','编码为空':'FFF2CC','待确认':'FFF2CC'}
    for name, rows, table_name in [('需处理',pending,'SimplePending'),('全部结果',all_rows,'SimpleAll')]:
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet['A2'] = f'{name}（{len(rows)}条）'
        sheet['A2'].font = Font(name='Microsoft YaHei',size=14,bold=True,color='243D5B')
        sheet['A3'] = '候选包含编码及ERP名称；按明示规格证据排序，证据相同时并列，仅供核对参考。'
        sheet['A3'].font = Font(name='Microsoft YaHei',size=10,color='64748B')
        for column, header in enumerate(HEADERS,1):
            cell = sheet.cell(4,column,header)
            cell.font = Font(name='Microsoft YaHei',size=10,bold=True,color='FFFFFF')
            cell.fill = PatternFill('solid',fgColor='243D5B')
            cell.alignment = Alignment(horizontal='center',vertical='center')
            sheet.column_dimensions[get_column_letter(column)].width = widths[column-1]
        sheet.row_dimensions[4].height = 30
        for number, row in enumerate(rows,5):
            lines = 2
            for column, header in enumerate(HEADERS,1):
                value = str(row[header] if row[header] is not None else '')
                cell = sheet.cell(number,column,value)
                cell.data_type = 's'
                cell.font = Font(name='Microsoft YaHei',size=10,color='243447')
                cell.alignment = Alignment(vertical='center',wrap_text=True)
                if column in (2,3,6,8,10,11,12):
                    cell.number_format = '@'
                cell_lines = sum(max(1,(sum(2 if ord(c)>255 else 1 for c in part)+widths[column-1]-3)//(widths[column-1]-2)) for part in value.split('\n'))
                lines = max(lines,cell_lines)
            if row['核验结果'] in fills:
                sheet.cell(number,7).fill = PatternFill('solid',fgColor=fills[row['核验结果']])
            sheet.row_dimensions[number].height = min(409,lines*14+8)
        if rows:
            table = Table(displayName=table_name,ref=f'A4:L{len(rows)+4}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            sheet.add_table(table)
        sheet.freeze_panes = 'D5'
    book.active = 0
    Path(output_path).parent.mkdir(parents=True,exist_ok=True)
    book.save(output_path)
    book.close()
